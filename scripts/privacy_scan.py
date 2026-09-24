import argparse
import csv
import io
import json
import os
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")

import django  # noqa: E402

django.setup()

from dashboard.insights import validate_content  # noqa: E402
from dashboard.models import DailySnapshot, ManagementInsightRevision  # noqa: E402
from interest_engine.analysis import (  # noqa: E402
    DETAIL_FIELDS as INTEREST_DETAIL_FIELDS,
    LEGACY_DETAIL_FIELDS as INTEREST_LEGACY_DETAIL_FIELDS,
    LEGACY_PRIVATE_DETAIL_FIELDS as INTEREST_LEGACY_PRIVATE_DETAIL_FIELDS,
    MANIFEST_SCHEMA as INTEREST_MANIFEST_SCHEMA,
    PREVIOUS_PRIVATE_DETAIL_FIELDS as INTEREST_PREVIOUS_PRIVATE_DETAIL_FIELDS,
    PRIVATE_DETAIL_FIELDS as INTEREST_PRIVATE_DETAIL_FIELDS,
)
from interest_engine.batch import BATCH_RESULT_SCHEMA  # noqa: E402
from interest_engine.candidates import CONTACT_LIKE  # noqa: E402
from interest_engine.contracts import sha256_file  # noqa: E402
from interest_engine.verify import verify as verify_interest_output  # noqa: E402
from interests.api import serialize_candidate, serialize_entity  # noqa: E402
from interests.models import InterestCandidate, InterestEntity, InterestSegmentSnapshot, InterestSnapshot  # noqa: E402
from publisher.privacy import EMAIL_VALUE, PHONE_VALUE, assert_aggregate_only  # noqa: E402
from publisher.interest_snapshot import validate_interest_payload  # noqa: E402
from publisher.snapshot import validate_snapshot_payload  # noqa: E402
from topics.models import TopicArtifact, TopicJob  # noqa: E402
from topics.storage import job_directory, safe_job_path  # noqa: E402


parser = argparse.ArgumentParser()
parser.add_argument("--interest-output", action="append", default=[], type=Path)
parser.add_argument("--interest-batch", action="append", default=[], type=Path)
parser.add_argument(
    "--registry",
    type=Path,
    default=PROJECT_ROOT / "interest_engine" / "resources" / "entity-registry-v1.json",
)
args = parser.parse_args()


def load_json(path: Path) -> dict:
    payload = json.loads(path.read_text(encoding="utf-8"))
    assert_aggregate_only(payload)
    return payload


def assert_safe_text(value: str, label: str) -> None:
    if EMAIL_VALUE.search(value) or PHONE_VALUE.search(value):
        raise SystemExit(f"Contact-like value found in {label}")


def scan_interest_detail_archive(archive_path: Path) -> int:
    rows = 0
    with zipfile.ZipFile(archive_path) as archive:
        names = archive.namelist()
        if len(names) != 1 or not names[0].endswith(".csv"):
            raise SystemExit("Interest detail ZIP must contain exactly one CSV")
        with archive.open(names[0]) as raw_stream:
            stream = io.TextIOWrapper(raw_stream, encoding="utf-8-sig", newline="")
            reader = csv.DictReader(stream)
            header = tuple(reader.fieldnames or ())
            if header not in {INTEREST_DETAIL_FIELDS, INTEREST_LEGACY_DETAIL_FIELDS}:
                raise SystemExit("Interest detail ZIP contains an unsafe header")
            for row in reader:
                rows += 1
                if re.fullmatch(r"[0-9a-f]{20}", row["user_key"] or "") is None:
                    raise SystemExit("Interest detail contains an invalid user pseudonym")
                if re.fullmatch(r"[0-9a-f]{20}", row["session_key"] or "") is None:
                    raise SystemExit("Interest detail contains an invalid session pseudonym")
                for key in ("route", "broad_topic", "behavior", "entity_names", "entity_ids", "interest_signal"):
                    if key not in header:
                        continue
                    assert_safe_text(row[key] or "", "interest detail")
    return rows


def scan_private_interest_detail(path: Path) -> int:
    rows = 0
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if tuple(reader.fieldnames or ()) not in {
            INTEREST_PRIVATE_DETAIL_FIELDS,
            INTEREST_PREVIOUS_PRIVATE_DETAIL_FIELDS,
            INTEREST_LEGACY_PRIVATE_DETAIL_FIELDS,
        }:
            raise SystemExit("Private interest detail contains an invalid header")
        for row in reader:
            rows += 1
            if not row.get("source_row"):
                raise SystemExit("Private interest detail contains a row without source_row")
    return rows


def scan_interest_workbook(workbook_path: Path) -> int:
    cells = 0
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    cells += 1
                    if isinstance(value, str):
                        assert_safe_text(value, "interest workbook")
    finally:
        workbook.close()
    return cells


def scan_private_candidates(path: Path) -> int:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    rows = payload.get("candidates", [])
    if not isinstance(rows, list) or len(rows) > 200:
        raise SystemExit("Private candidate package exceeds the 200-item limit")
    for row in rows:
        blocked = {"clientId", "cid", "user_key", "session_key", "timestamp", "response", "ai_text"}.intersection(row)
        if blocked:
            raise SystemExit(f"Private candidate package contains identifiers: {', '.join(sorted(blocked))}")
        examples = row.get("examples", [])
        if not isinstance(examples, list) or len(examples) > 3:
            raise SystemExit("Private candidate package exceeds three contexts")
        for example in examples:
            if CONTACT_LIKE.search(str(example)):
                raise SystemExit("Private candidate context is not de-identified")
    return len(rows)


seed_files = sorted((PROJECT_ROOT / "dashboard" / "data" / "seed_snapshots").glob("*.json"))
if not seed_files:
    raise SystemExit("No seed snapshots found")

seed_payloads = []
for seed_file in seed_files:
    payload = load_json(seed_file)
    validate_snapshot_payload(payload)
    seed_payloads.append(payload)

baseline = load_json(PROJECT_ROOT / "dashboard" / "data" / "historical_baseline.json")
database_payloads = list(DailySnapshot.objects.order_by("data_date").values_list("payload", flat=True))
for payload in database_payloads:
    validate_snapshot_payload(payload)
    assert_aggregate_only(payload)

api_shape = {
    "schemaVersion": "teeni-dashboard-response/1.1.0",
    "snapshots": database_payloads,
    "baseline": baseline,
    "latestDate": database_payloads[-1]["dataDate"] if database_payloads else None,
}
assert_aggregate_only(api_shape)

insight_revisions = list(ManagementInsightRevision.objects.order_by("insight_id", "revision"))
for revision in insight_revisions:
    validate_content(
        {
            "title": revision.title,
            "summary": revision.summary,
            "facts": revision.facts,
            "mechanisms": revision.mechanisms,
            "associations": revision.associations,
            "hypotheses": revision.hypotheses,
            "actions": revision.actions,
        }
    )
    assert_aggregate_only(revision.metric_snapshot)
    if not isinstance(revision.source_hashes, dict) or any(
        re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(key)) is None
        or re.fullmatch(r"[0-9a-f]{64}", str(value)) is None
        for key, value in revision.source_hashes.items()
    ):
        raise SystemExit("Management insight contains invalid source hashes")

interest_snapshots = list(InterestSnapshot.objects.select_related("job").order_by("data_date"))
for snapshot in interest_snapshots:
    validate_interest_payload(snapshot.payload, snapshot.job)

interest_segment_snapshots = list(InterestSegmentSnapshot.objects.order_by("data_date", "dimension_set"))
for segment in interest_segment_snapshots:
    if segment.schema_version != "teeni-interest-segments/1.0.0":
        raise SystemExit("Interest segment database schema mismatch")
    if segment.payload.get("dimensionSet") != segment.dimension_set:
        raise SystemExit("Interest segment database dimension mismatch")
    assert_aggregate_only(segment.payload)

candidate_payloads = [
    serialize_candidate(candidate)
    for candidate in InterestCandidate.objects.select_related("resolved_entity").order_by("data_date", "id")
]
entity_payloads = [
    serialize_entity(entity)
    for entity in InterestEntity.objects.prefetch_related("aliases").order_by("registry_order", "id")
]
assert_aggregate_only({"candidates": candidate_payloads, "entities": entity_payloads})

interest_detail_rows = 0
interest_workbook_cells = 0
interest_artifacts = TopicArtifact.objects.filter(
    verified=True,
    kind__in=[TopicArtifact.Kind.INTEREST_WORKBOOK, TopicArtifact.Kind.INTEREST_DETAIL_ZIP],
)
artifact_paths = set()
for artifact in interest_artifacts:
    path = safe_job_path(artifact.job_id, artifact.relative_path)
    if not path.is_file():
        continue
    artifact_paths.add(path.resolve())
    if artifact.kind == TopicArtifact.Kind.INTEREST_DETAIL_ZIP:
        interest_detail_rows += scan_interest_detail_archive(path)
    else:
        interest_workbook_cells += scan_interest_workbook(path)

private_candidate_files = []
private_interest_detail_files = []
private_interest_detail_rows = 0
for job in TopicJob.objects.filter(
    pipeline__in=[TopicJob.Pipeline.SHADOW, TopicJob.Pipeline.INTEREST_V1]
):
    assert_aggregate_only(job.runtime_state)
    root = job_directory(job.id)
    if not root.is_dir():
        continue
    if any(root.rglob(".interest-work.private.sqlite3")):
        raise SystemExit("Interest analysis left a private SQLite work file behind")
    for path in root.rglob("teeni-interest-candidates.private.json"):
        if path.resolve() in artifact_paths or "artifacts" in path.parts:
            raise SystemExit("Private candidate evidence was exposed as a download artifact")
        scan_private_candidates(path)
        private_candidate_files.append(path)
    private_value = job.runtime_state.get("interest", {}).get("privateDetailPath")
    for path in [safe_job_path(job.id, private_value)] if private_value else []:
        if path.resolve() in artifact_paths or "artifacts" in path.parts:
            raise SystemExit("Private original-text detail was exposed as a download artifact")
        private_interest_detail_rows += scan_private_interest_detail(path)
        private_interest_detail_files.append(path)

static_files = sorted((PROJECT_ROOT / "static" / "dashboard").rglob("*"))
static_files = [path for path in static_files if path.is_file()]
for static_file in static_files:
    text = static_file.read_text(encoding="utf-8")
    assert_safe_text(text, f"built frontend: {static_file.name}")

log_files = [path for path in (PROJECT_ROOT / "logs").rglob("*") if path.is_file()]
serialized_private_field = re.compile(r'["\'](?:clientId|cid|ai_text|query|response|examples)["\']\s*:', re.IGNORECASE)
for log_file in log_files:
    text = log_file.read_text(encoding="utf-8", errors="replace")
    assert_safe_text(text, f"log: {log_file.name}")
    if serialized_private_field.search(text):
        raise SystemExit(f"Serialized private field found in log: {log_file.name}")

verified_external_outputs = 0
external_private_candidates = 0
external_private_detail_rows = 0


def scan_external_interest_output(output_dir: Path, base_detail_path: Path | None = None) -> tuple[int, int]:
    output_dir = output_dir.resolve()
    manifest_path = output_dir / "teeni-interest-manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
    assert_aggregate_only(manifest)
    current_contract = manifest.get("schemaVersion") == INTEREST_MANIFEST_SCHEMA
    safe_detail_path = output_dir / (
        "teeni-interest-detail.safe.csv" if current_contract else "teeni-interest-detail.csv"
    )
    private_detail_path = output_dir / (
        "teeni-interest-detail.csv" if current_contract else "teeni-interest-detail.private.csv"
    )
    segment_path = output_dir / "teeni-interest-segments.json" if current_contract else None
    verify_interest_output(
        output_dir / "teeni-interest-report.xlsx",
        safe_detail_path,
        output_dir / "teeni-interest-snapshot.json",
        output_dir / "teeni-interest-candidates.private.json",
        manifest_path,
        args.registry,
        private_detail_path if private_detail_path.is_file() else None,
        base_detail_path,
        segment_path,
    )
    scan_interest_workbook(output_dir / "teeni-interest-report.xlsx")
    private_candidates = scan_private_candidates(output_dir / "teeni-interest-candidates.private.json")
    private_detail_rows = 0
    if private_detail_path.is_file():
        private_detail_rows = scan_private_interest_detail(private_detail_path)
    if any(output_dir.rglob(".interest-work.private.sqlite3")):
        raise SystemExit("External interest output left a private SQLite work file behind")
    snapshot_payload = load_json(output_dir / "teeni-interest-snapshot.json")
    if segment_path is not None:
        segment_payload = load_json(segment_path)
        if segment_payload.get("schemaVersion") != "teeni-interest-segments/1.0.0":
            raise SystemExit("Interest segment sidecar schema mismatch")
        if len(segment_payload.get("dimensionSets", {})) != 7:
            raise SystemExit("Interest segment sidecar must contain seven dimension sets")
    if snapshot_payload.get("restrictedEntities") is None:
        raise SystemExit("Interest snapshot is missing restricted aggregate isolation")
    return private_candidates, private_detail_rows


for output_dir in args.interest_output:
    private_candidates, private_detail_rows = scan_external_interest_output(output_dir)
    external_private_candidates += private_candidates
    external_private_detail_rows += private_detail_rows
    verified_external_outputs += 1

verified_interest_batches = 0
for batch_dir in args.interest_batch:
    batch_dir = batch_dir.resolve()
    result_path = batch_dir / "teeni-interest-batch-result.json"
    result = load_json(result_path)
    if result.get("schemaVersion") != BATCH_RESULT_SCHEMA:
        raise SystemExit("Interest batch result schema mismatch")
    dates = result.get("dates")
    if not isinstance(dates, list) or not 1 <= len(dates) <= 7 or len(set(dates)) != len(dates):
        raise SystemExit("Interest batch result must contain one to seven unique dates")
    daily = result.get("daily")
    if not isinstance(daily, list) or [row.get("dataDate") for row in daily] != dates:
        raise SystemExit("Interest batch dates do not match daily outputs")

    def batch_path(value, *, artifact=True):
        if not isinstance(value, str) or not value:
            raise SystemExit("Interest batch path is missing")
        path = Path(value)
        resolved = (batch_dir / path).resolve() if not path.is_absolute() else path.resolve()
        if artifact and not resolved.is_relative_to(batch_dir):
            raise SystemExit("Interest batch artifact escapes result directory")
        return resolved

    for row in daily:
        snapshot_path = batch_path(row.get("snapshotPath"))
        manifest_path = batch_path(row.get("manifestPath"))
        segment_path = batch_path(row.get("segmentPath"))
        base_detail_path = batch_path(row.get("baseDetailPath"), artifact=False)
        if sha256_file(snapshot_path) != row.get("snapshotSha256"):
            raise SystemExit("Interest batch snapshot SHA-256 mismatch")
        if sha256_file(manifest_path) != row.get("manifestSha256"):
            raise SystemExit("Interest batch manifest SHA-256 mismatch")
        if sha256_file(segment_path) != row.get("segmentSha256"):
            raise SystemExit("Interest batch segment SHA-256 mismatch")
        if sha256_file(base_detail_path) != row.get("baseDetailSha256"):
            raise SystemExit("Interest batch base detail SHA-256 mismatch")
        private_candidates, private_detail_rows = scan_external_interest_output(
            snapshot_path.parent,
            base_detail_path,
        )
        external_private_candidates += private_candidates
        external_private_detail_rows += private_detail_rows
        verified_external_outputs += 1
    candidate_path = batch_path(result.get("candidatePath"))
    if sha256_file(candidate_path) != result.get("candidateSha256"):
        raise SystemExit("Interest batch candidate SHA-256 mismatch")
    batch_candidates = scan_private_candidates(candidate_path)
    batch_payload = json.loads(candidate_path.read_text(encoding="utf-8-sig"))
    if batch_payload.get("candidateBatchSize") != 20 or batch_payload.get("candidateRequestLimit") != 10:
        raise SystemExit("Interest batch candidate request limits mismatch")
    if int(batch_payload.get("candidateBatches", 0)) > 10:
        raise SystemExit("Interest batch candidate request count exceeds ten")
    external_private_candidates += batch_candidates
    verified_interest_batches += 1

print(
    json.dumps(
        {
            "seedSnapshots": len(seed_payloads),
            "databaseSnapshots": len(database_payloads),
            "managementInsightRevisions": len(insight_revisions),
            "interestSnapshots": len(interest_snapshots),
            "interestSegmentSnapshots": len(interest_segment_snapshots),
            "interestCandidates": len(candidate_payloads),
            "interestEntities": len(entity_payloads),
            "interestDetailRows": interest_detail_rows,
            "interestWorkbookCells": interest_workbook_cells,
            "privateCandidateFiles": len(private_candidate_files),
            "privateInterestDetailFiles": len(private_interest_detail_files),
            "privateInterestDetailRows": private_interest_detail_rows,
            "externalInterestOutputs": verified_external_outputs,
            "verifiedInterestBatches": verified_interest_batches,
            "externalPrivateCandidates": external_private_candidates,
            "externalPrivateDetailRows": external_private_detail_rows,
            "latestDate": api_shape["latestDate"],
            "builtFrontendFiles": len(static_files),
            "logFiles": len(log_files),
            "status": "aggregate-only",
        },
        ensure_ascii=False,
    )
)
