from __future__ import annotations

import csv
import hashlib
import json
import re
import sqlite3
import tempfile
from collections import Counter
from dataclasses import dataclass
from itertools import zip_longest
from pathlib import Path

from openpyxl import load_workbook

from . import ENGINE_VERSION, REPORT_SCHEMA, SNAPSHOT_SCHEMA
from .analysis import (
    CANDIDATE_SCHEMA,
    DETAIL_FIELDS,
    DETAIL_SCHEMA,
    EARLIER_MANIFEST_SCHEMA,
    INITIAL_MANIFEST_SCHEMA,
    LEGACY_DETAIL_FIELDS,
    LEGACY_DETAIL_SCHEMA,
    LEGACY_MANIFEST_SCHEMA,
    LEGACY_PRIVATE_DETAIL_FIELDS,
    LEGACY_PRIVATE_DETAIL_SCHEMA,
    MANIFEST_SCHEMA,
    OLDER_MANIFEST_SCHEMA,
    PREVIOUS_MANIFEST_SCHEMA,
    PREVIOUS_PRIVATE_DETAIL_FIELDS,
    PREVIOUS_PRIVATE_DETAIL_SCHEMA,
    PRIVATE_DETAIL_FIELDS,
    PRIVATE_DETAIL_SCHEMA,
)
from .candidates import CONTACT_LIKE
from .classifier import Behavior, InterestSignal, QueryRoute, Turn, TurnContext, classify_turn, detail_evidence
from .contracts import REQUIRED_DETAIL_FIELDS
from .contracts import SCENE_PRODUCTS, product_identity, sha256_file
from .registry import BroadTopic, EntityRegistry, EntityType, Visibility, is_standalone_ip
from .segments import (
    DEFAULT_PROVINCE_MAP,
    DIMENSION_SETS,
    GENDERS as SEGMENT_GENDERS,
    SEGMENT_SCHEMA,
    ProvinceMap,
    build_segment_payload,
    legacy_demographic,
    materialize_group,
    validate_segment_payload,
)


EXPECTED_SHEETS = (
    "兴趣结论总览",
    "IP总览",
    "画像分层与IP",
    "实体明细",
    "风险热梗",
    "新梗雷达",
    "普通话题",
    "用户行为",
    "清洗与路由",
    "实体词典",
    "分析口径",
)
FORBIDDEN_PUBLIC_KEYS = {
    "query",
    "text",
    "ai_text",
    "clientId",
    "cid",
    "timestamp",
    "created_at",
    "response",
    "context",
    "examples",
    "user_key",
    "session_key",
    "source_row",
    "profile_age",
    "profile_gender",
    "age_status",
    "gender_status",
    "city",
    "city_normalized",
    "city_status",
    "region_id",
    "province",
    "birthday",
}
AGES = tuple(range(1, 18))
GENDERS = ("男", "女")
DEMOGRAPHIC_SAMPLE_MIN_USERS = 30


class VerificationError(ValueError):
    pass


@dataclass(frozen=True)
class VerificationResult:
    ok: bool
    detail_rows: int
    output_count: int
    sheet_count: int
    private_detail_rows: int = 0


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise VerificationError(message)


def _load_json(path: Path, label: str) -> dict:
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise VerificationError(f"{label} is not valid UTF-8 JSON") from exc
    _require(isinstance(value, dict), f"{label} must be a JSON object")
    return value


def _assert_public_keys(value) -> None:
    if isinstance(value, dict):
        blocked = FORBIDDEN_PUBLIC_KEYS.intersection(value)
        _require(not blocked, f"aggregate snapshot contains forbidden keys: {', '.join(sorted(blocked))}")
        for child in value.values():
            _assert_public_keys(child)
    elif isinstance(value, list):
        for child in value:
            _assert_public_keys(child)


def _recompute_detail(detail_path: Path, registry: EntityRegistry, detail_fields: tuple[str, ...]) -> dict:
    temporary = tempfile.NamedTemporaryFile(prefix="teeni-interest-verify-", suffix=".sqlite3", delete=False)
    temporary.close()
    database_path = Path(temporary.name)
    database = sqlite3.connect(database_path)
    routes: Counter[str] = Counter()
    topics: Counter[str] = Counter()
    behaviors: Counter[str] = Counter()
    row_count = 0
    valid_content = 0
    entity_queries = 0
    try:
        database.execute(
            "CREATE TABLE entity_turn(entity_id TEXT, user_key TEXT, session_key TEXT, query_key TEXT, signal TEXT, polarity TEXT)"
        )
        with detail_path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            _require(tuple(reader.fieldnames or ()) == detail_fields, "interest detail header is invalid")
            entity_lookup = {entity.id: entity.canonical_name for entity in registry.entities}
            for row in reader:
                row_count += 1
                _require(re.fullmatch(r"[0-9a-f]{20}", row["user_key"] or "") is not None, "invalid user pseudonym")
                _require(re.fullmatch(r"[0-9a-f]{20}", row["session_key"] or "") is not None, "invalid session pseudonym")
                route = QueryRoute(row["route"])
                signal = InterestSignal(row["interest_signal"])
                routes[route.value] += 1
                entity_ids = tuple(value for value in row["entity_ids"].split("|") if value)
                entity_signals = tuple(filter(None, row.get("entity_signals", "").split("|")))
                polarities = tuple(filter(None, row.get("preference_polarities", "").split("|")))
                bases = tuple(filter(None, row.get("match_bases", "").split("|")))
                _require(len(entity_ids) == len(entity_signals) == len(polarities) == len(bases), "entity evidence alignment mismatch")
                _require(all(value in {item.value for item in InterestSignal} for value in entity_signals), "invalid entity signal")
                _require(all(value in {"positive", "negative", "neutral", "mixed"} for value in polarities), "invalid preference polarity")
                _require(all(value in {"single_context", "explicit_alias"} for value in bases), "invalid match basis")
                entity_names = tuple(value for value in row.get("entity_names", "").split("|") if value)
                _require(len(entity_ids) == len(set(entity_ids)), "detail has duplicate entity ids")
                if "entity_names" in detail_fields:
                    _require(
                        entity_names == tuple(entity_lookup.get(entity_id) for entity_id in entity_ids),
                        "detail entity names do not match the registry",
                    )
                if route == QueryRoute.VALID_CONTENT:
                    valid_content += 1
                    topic = BroadTopic(row["broad_topic"])
                    behavior = Behavior(row["behavior"])
                    topics[topic.value] += 1
                    behaviors[behavior.value] += 1
                else:
                    _require(
                        not row["broad_topic"] and not entity_ids and not entity_names,
                        "non-content row entered content analysis",
                    )
                if entity_ids:
                    entity_queries += 1
                    _require(route == QueryRoute.VALID_CONTENT, "entity match exists outside valid content")
                    database.executemany(
                        "INSERT INTO entity_turn VALUES (?, ?, ?, ?, ?, ?)",
                        (
                            (entity_id, row["user_key"], row["session_key"], str(row_count), entity_signal, polarity)
                            for entity_id, entity_signal, polarity in zip(entity_ids, entity_signals, polarities, strict=True)
                        ),
                    )
        database.commit()
        known = {entity.id for entity in registry.entities}
        actual = {row[0] for row in database.execute("SELECT DISTINCT entity_id FROM entity_turn")}
        _require(actual.issubset(known), "detail references an unknown entity")
        def metrics(entity_ids: tuple[str, ...]) -> dict:
            placeholders = ",".join("?" for _ in entity_ids)
            queries, mention_users, sessions, initiators, continuers = database.execute(
                """
                SELECT COUNT(DISTINCT query_key), COUNT(DISTINCT user_key), COUNT(DISTINCT session_key),
                       COUNT(DISTINCT CASE WHEN signal = ? THEN user_key END),
                       COUNT(DISTINCT CASE WHEN signal = ? THEN user_key END)
                FROM entity_turn WHERE entity_id IN (%s)
                """ % placeholders,
                (InterestSignal.INITIATION.value, InterestSignal.CONTINUATION.value, *entity_ids),
            ).fetchone()
            active_users = database.execute(
                f"SELECT COUNT(DISTINCT user_key) FROM entity_turn WHERE entity_id IN ({placeholders}) AND signal IN (?, ?)",
                (*entity_ids, InterestSignal.INITIATION.value, InterestSignal.CONTINUATION.value),
            ).fetchone()[0]
            deep_sessions = database.execute(
                f"""
                SELECT COUNT(*) FROM (
                    SELECT session_key FROM entity_turn WHERE entity_id IN ({placeholders})
                    GROUP BY session_key HAVING COUNT(DISTINCT query_key) >= 3
                )
                """,
                entity_ids,
            ).fetchone()[0]
            polarities = {polarity: database.execute(
                f"SELECT COUNT(DISTINCT user_key) FROM entity_turn WHERE entity_id IN ({placeholders}) AND polarity = ?",
                (*entity_ids, polarity),
            ).fetchone()[0] for polarity in ("positive", "negative", "neutral", "mixed")}
            return {
                "positivePreferenceUsers": polarities["positive"],
                "negativePreferenceUsers": polarities["negative"],
                "neutralMentionUsers": polarities["neutral"],
                "mixedPreferenceUsers": polarities["mixed"],
                "sampleUsers": mention_users,
                "sampleStatus": "可描述" if mention_users >= 30 else "样本不足",
                "deepSampleSessions": sessions,
                "activeInterestUsers": active_users,
                "mentionUsers": mention_users,
                "initiatorUsers": initiators,
                "continuationUsers": continuers,
                "sessions": sessions,
                "queries": queries,
                "deepSessions": deep_sessions,
                "deepChatRate": deep_sessions / sessions if sessions else 0,
            }

        public_actual = {entity_id for entity_id in actual if registry.get(entity_id).visibility == Visibility.PUBLIC}
        restricted_actual = actual - public_actual
        entities = {entity_id: metrics((entity_id,)) for entity_id in public_actual}
        restricted_entities = {entity_id: metrics((entity_id,)) for entity_id in restricted_actual}
        rollups: dict[str, dict] = {}
        parent_ids = {entity.parent_registry_id for entity in registry.entities if entity.parent_registry_id}
        standalone_ids = {
            entity_id for entity_id in public_actual if is_standalone_ip(registry.get(entity_id))
        }
        for rollup_id in {parent_id for parent_id in parent_ids if parent_id} | standalone_ids:
            rollup = registry.get(rollup_id)
            if rollup.visibility != Visibility.PUBLIC:
                continue
            is_parent = rollup_id in parent_ids
            member_ids = (
                tuple(
                    entity_id
                    for entity_id in public_actual
                    if entity_id == rollup_id or rollup_id in {ancestor.id for ancestor in registry.ancestors(entity_id)}
                )
                if is_parent
                else (rollup_id,)
            )
            if member_ids:
                rollups[rollup_id] = {
                    **metrics(member_ids),
                    "childEntityCount": len([entity_id for entity_id in member_ids if entity_id != rollup_id]),
                    "rollupKind": "parent" if is_parent else "standalone",
                }
        structures: dict[str, dict] = {}
        structure_ids: dict[str, list[str]] = {}
        for entity_id in public_actual:
            entity = registry.get(entity_id)
            name = f"{entity.entity_type.value} / {entity.entity_subtype.value}"
            structure_ids.setdefault(name, []).append(entity_id)
        for name, entity_ids in structure_ids.items():
            structures[name] = metrics(tuple(entity_ids))
        if public_actual:
            placeholders = ",".join("?" for _ in public_actual)
            public_ids = tuple(public_actual)
            entity_queries = database.execute(
                f"SELECT COUNT(DISTINCT query_key) FROM entity_turn WHERE entity_id IN ({placeholders})",
                public_ids,
            ).fetchone()[0]
            active_users = database.execute(
                f"""
                SELECT COUNT(DISTINCT user_key) FROM entity_turn
                WHERE entity_id IN ({placeholders}) AND signal IN (?, ?)
                """,
                (*public_ids, InterestSignal.INITIATION.value, InterestSignal.CONTINUATION.value),
            ).fetchone()[0]
        else:
            entity_queries = 0
            active_users = 0
        return {
            "rows": row_count,
            "valid": valid_content,
            "entityQueries": entity_queries,
            "activeUsers": active_users,
            "routes": routes,
            "topics": topics,
            "behaviors": behaviors,
            "entities": entities,
            "restrictedEntities": restricted_entities,
            "ipRollups": rollups,
            "entityStructures": structures,
        }
    finally:
        database.close()
        database_path.unlink(missing_ok=True)


def _verify_private_detail(
    private_detail_path: Path,
    detail_path: Path,
    private_fields: tuple[str, ...],
    detail_fields: tuple[str, ...],
) -> int:
    compared_fields = tuple(
        field
        for field in ("source_row", "route", "broad_topic", "behavior", "entity_names", "entity_ids", "interest_signal", "entity_signals", "preference_polarities", "match_bases")
        if field in detail_fields and field in private_fields
    )
    rows = 0
    with (
        private_detail_path.open("r", encoding="utf-8-sig", newline="") as private_stream,
        detail_path.open("r", encoding="utf-8-sig", newline="") as safe_stream,
    ):
        private_reader = csv.DictReader(private_stream)
        safe_reader = csv.DictReader(safe_stream)
        _require(tuple(private_reader.fieldnames or ()) == private_fields, "private detail header is invalid")
        _require(tuple(safe_reader.fieldnames or ()) == detail_fields, "interest detail header is invalid")
        for private_row, safe_row in zip_longest(private_reader, safe_reader):
            _require(private_row is not None and safe_row is not None, "private and safe detail row counts differ")
            rows += 1
            for field in compared_fields:
                _require(private_row[field] == safe_row[field], f"private detail {field} mismatch")
    return rows


def _verification_key(prefix: str, value: str, detail_hash: str) -> str:
    return hashlib.sha256(f"{prefix}\0{detail_hash}\0{value}".encode("utf-8")).hexdigest()[:20]


def _verified_age(value: object) -> int | None:
    try:
        age = int(str(value).strip())
    except (TypeError, ValueError):
        return None
    return age if age in AGES else None


def _verified_dimensions(row: dict[str, str], client_id: str, province_map: ProvinceMap) -> dict[str, object]:
    if not client_id:
        return {
            "age": None, "age_stable": 0,
            "gender": None, "gender_stable": 0,
            "region_id": None, "region_stable": 0, "region_unmapped": 0,
        }
    age = _verified_age(row.get("profile_age")) if row.get("age_status") == "正常" else None
    gender_label = str(row.get("profile_gender") or "").strip()
    gender = {"男": "male", "女": "female"}.get(gender_label) if row.get("gender_status") == "正常" else None
    city = str(row.get("city_normalized") or "").strip()
    city_normal = row.get("city_status") == "正常" and bool(city)
    region_id = province_map.region_for_city(city) if city_normal else None
    return {
        "age": age,
        "age_stable": 1 if age is not None else 0,
        "gender": gender,
        "gender_stable": 1 if gender is not None else 0,
        "region_id": region_id,
        "region_stable": 1 if region_id is not None else 0,
        "region_unmapped": 1 if city_normal and region_id is None else 0,
    }


def _recompute_segments(
    base_detail_path: Path,
    detail_path: Path,
    registry: EntityRegistry,
    detail_hash: str,
    data_date: str,
    registry_hash: str,
    province_map: ProvinceMap,
    scene_id: str = "488",
) -> dict:
    temporary = tempfile.NamedTemporaryFile(prefix="teeni-interest-demographic-verify-", suffix=".sqlite3", delete=False)
    temporary.close()
    database_path = Path(temporary.name)
    database = sqlite3.connect(database_path)
    try:
        database.executescript(
            """
            CREATE TABLE session_state (
                session_key TEXT PRIMARY KEY,
                user_entities TEXT NOT NULL,
                ai_entities TEXT NOT NULL
            );
            CREATE TABLE entity_turn (
                entity_id TEXT NOT NULL,
                user_key TEXT NOT NULL,
                query_key TEXT NOT NULL,
                signal TEXT NOT NULL
            );
            CREATE TABLE user_demographic (
                user_key TEXT PRIMARY KEY,
                age INTEGER,
                age_stable INTEGER NOT NULL,
                gender TEXT,
                gender_stable INTEGER NOT NULL,
                region_id TEXT,
                region_stable INTEGER NOT NULL,
                region_unmapped INTEGER NOT NULL,
                eligible INTEGER NOT NULL DEFAULT 0
            );
            """
        )
        seen_source_rows: set[str] = set()
        with (
            base_detail_path.open("r", encoding="utf-8-sig", newline="") as base_stream,
            detail_path.open("r", encoding="utf-8-sig", newline="") as safe_stream,
        ):
            base_reader = csv.DictReader(base_stream)
            safe_reader = csv.DictReader(safe_stream)
            _require(REQUIRED_DETAIL_FIELDS.issubset(set(base_reader.fieldnames or ())), "base detail demographic header is invalid")
            _require(tuple(safe_reader.fieldnames or ()) == DETAIL_FIELDS, "interest detail header is invalid")
            for base_row, safe_row in zip_longest(base_reader, safe_reader):
                _require(base_row is not None and safe_row is not None, "base and interest detail row counts differ")
                _require(str(base_row.get("sceneId", "")).strip() == scene_id,
                         "base detail sceneId mismatch or mixed scenes")
                source_row = str(base_row.get("source_row") or "")
                _require(re.fullmatch(r"[1-9][0-9]*", source_row) is not None, "base detail source_row is invalid")
                _require(source_row not in seen_source_rows, "base detail source_row is duplicated")
                seen_source_rows.add(source_row)
                _require(safe_row["source_row"] == source_row, "base and interest detail source_row mismatch")
                client_id = str(base_row.get("clientId") or "")
                cid = str(base_row.get("cid") or "")
                session_key = _verification_key("session", cid or f"row:{source_row}", detail_hash)
                user_key = _verification_key("user", client_id or f"session:{session_key}", detail_hash)
                _require(safe_row["session_key"] == session_key, "interest detail session pseudonym mismatch")
                _require(safe_row["user_key"] == user_key, "interest detail user pseudonym mismatch")

                previous = database.execute(
                    "SELECT user_entities, ai_entities FROM session_state WHERE session_key = ?",
                    (session_key,),
                ).fetchone()
                result = classify_turn(
                    Turn(
                        text=str(base_row.get("text") or ""),
                        ai_text=str(base_row.get("ai_text") or ""),
                        is_template=str(base_row.get("is_template") or "否"),
                        is_invalid_turn=str(base_row.get("is_invalid_turn") or "否"),
                        invalid_reason=str(base_row.get("invalid_reason") or ""),
                    ),
                    TurnContext(
                        previous_entity_ids=tuple(previous[0].split("|")) if previous and previous[0] else (),
                        previous_ai_entity_ids=tuple(previous[1].split("|")) if previous and previous[1] else (),
                    ),
                    registry,
                )
                entity_ids = tuple(match.entity.id for match in result.entity_matches)
                entity_names = tuple(match.entity.canonical_name for match in result.entity_matches)
                expected = {
                    **detail_evidence(result),
                    "route": result.route.value,
                    "broad_topic": result.broad_topic.value if result.broad_topic else "",
                    "behavior": result.behavior.value if result.behavior else "",
                    "entity_names": "|".join(entity_names),
                    "entity_ids": "|".join(entity_ids),
                    "interest_signal": result.interest_signal.value,
                }
                for field, value in expected.items():
                    _require(safe_row[field] == value, f"interest detail {field} does not match base detail")

                demographic = _verified_dimensions(base_row, client_id, province_map)
                database.execute(
                    """
                    INSERT INTO user_demographic(
                        user_key, age, age_stable, gender, gender_stable,
                        region_id, region_stable, region_unmapped, eligible
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, 0)
                    ON CONFLICT(user_key) DO UPDATE SET
                        age = CASE
                            WHEN user_demographic.age_stable = 1 AND excluded.age_stable = 1
                                 AND user_demographic.age = excluded.age
                            THEN user_demographic.age ELSE NULL END,
                        age_stable = CASE
                            WHEN user_demographic.age_stable = 1 AND excluded.age_stable = 1
                                 AND user_demographic.age = excluded.age
                            THEN 1 ELSE 0 END,
                        gender = CASE
                            WHEN user_demographic.gender_stable = 1 AND excluded.gender_stable = 1
                                 AND user_demographic.gender = excluded.gender
                            THEN user_demographic.gender ELSE NULL END,
                        gender_stable = CASE
                            WHEN user_demographic.gender_stable = 1 AND excluded.gender_stable = 1
                                 AND user_demographic.gender = excluded.gender
                            THEN 1 ELSE 0 END,
                        region_id = CASE
                            WHEN user_demographic.region_stable = 1 AND excluded.region_stable = 1
                                 AND user_demographic.region_id = excluded.region_id
                            THEN user_demographic.region_id ELSE NULL END,
                        region_stable = CASE
                            WHEN user_demographic.region_stable = 1 AND excluded.region_stable = 1
                                 AND user_demographic.region_id = excluded.region_id
                            THEN 1 ELSE 0 END,
                        region_unmapped = MAX(user_demographic.region_unmapped, excluded.region_unmapped)
                    """,
                    (
                        user_key,
                        demographic["age"],
                        demographic["age_stable"],
                        demographic["gender"],
                        demographic["gender_stable"],
                        demographic["region_id"],
                        demographic["region_stable"],
                        demographic["region_unmapped"],
                    ),
                )
                if result.route == QueryRoute.VALID_CONTENT:
                    database.execute("UPDATE user_demographic SET eligible = 1 WHERE user_key = ?", (user_key,))
                    database.executemany(
                        "INSERT INTO entity_turn(entity_id, user_key, query_key, signal) VALUES (?, ?, ?, ?)",
                        ((entity_id, user_key, source_row, signal.value) for entity_id, signal in zip(entity_ids, result.entity_signals, strict=True)),
                    )
                ai_entities = tuple(match.entity.id for match in registry.match(str(base_row.get("ai_text") or "")))
                database.execute(
                    """
                    INSERT INTO session_state(session_key, user_entities, ai_entities) VALUES (?, ?, ?)
                    ON CONFLICT(session_key) DO UPDATE SET user_entities = excluded.user_entities, ai_entities = excluded.ai_entities
                    """,
                    (
                        session_key,
                        "|".join(entity_ids) if result.route == QueryRoute.VALID_CONTENT else "",
                        "|".join(ai_entities),
                    ),
                )
        database.commit()

        matched = {row[0] for row in database.execute("SELECT DISTINCT entity_id FROM entity_turn")}
        parent_ids = {entity.parent_registry_id for entity in registry.entities if entity.parent_registry_id}
        standalone_ids = {
            entity.id for entity in registry.entities
            if entity.id in matched and entity.visibility == Visibility.PUBLIC and is_standalone_ip(entity)
        }
        rollup_members: dict[str, tuple[str, ...]] = {}
        for rollup_id in sorted({value for value in parent_ids if value} | standalone_ids):
            rollup = registry.get(rollup_id)
            if rollup.visibility != Visibility.PUBLIC:
                continue
            member_ids = (
                tuple(
                    entity.id for entity in registry.entities
                    if entity.id in matched
                    and entity.visibility == Visibility.PUBLIC
                    and (entity.id == rollup_id or rollup_id in {ancestor.id for ancestor in registry.ancestors(entity.id)})
                )
                if rollup_id in parent_ids else (rollup_id,)
            )
            if member_ids:
                rollup_members[rollup_id] = member_ids
        database.execute(
            "CREATE TABLE rollup_member(rollup_id TEXT NOT NULL, entity_id TEXT NOT NULL, PRIMARY KEY (rollup_id, entity_id))"
        )
        database.executemany(
            "INSERT INTO rollup_member(rollup_id, entity_id) VALUES (?, ?)",
            ((rollup_id, entity_id) for rollup_id, members in rollup_members.items() for entity_id in members),
        )
        return build_segment_payload(
            database,
            list(rollup_members),
            province_map,
            data_date=data_date,
            source_sha256=detail_hash,
            registry_sha256=registry_hash,
            product_version=SCENE_PRODUCTS[scene_id],
            scene_id=scene_id,
        )
    finally:
        database.close()
        database_path.unlink(missing_ok=True)


def _same_cell_value(actual, expected) -> bool:
    if isinstance(expected, float):
        return isinstance(actual, (int, float)) and abs(actual - expected) < 1e-12
    return actual == expected


def _require_sheet_row(sheet, row_number: int, expected: list, label: str) -> None:
    actual = [sheet.cell(row_number, column).value for column in range(1, len(expected) + 1)]
    _require(
        len(actual) == len(expected) and all(_same_cell_value(value, target) for value, target in zip(actual, expected)),
        f"workbook {label} row mismatch",
    )


def _verify_segment_sheet(sheet, snapshot: dict, segments: dict) -> None:
    names = {row["id"]: row["name"] for row in snapshot["ipRollups"]}
    gender_labels = dict(SEGMENT_GENDERS)
    region_labels = {row["id"]: row["label"] for row in segments["domains"]["regions"]}
    dimension_labels = {
        "age": "年龄",
        "gender": "性别",
        "region": "地区",
        "age_gender": "年龄×性别",
        "age_region": "年龄×地区",
        "gender_region": "性别×地区",
        "age_gender_region": "年龄×性别×地区",
    }
    _require(
        sheet.cell(1, 1).value == "七种画像口径各分组 Top 10 IP；pp = 组内覆盖率 - 当前维度口径总体覆盖率",
        "workbook segment Top 10 title mismatch",
    )
    top_headers = [
        "维度口径", "年龄", "性别", "地区", "分组用户", "兴趣可分析用户", "当前口径总体可分析用户",
        "排名", "IP", "兴趣用户", "组内覆盖率", "当前口径总体覆盖率", "百分点差（pp）", "样本状态",
    ]
    _require_sheet_row(sheet, 2, top_headers, "segment Top 10 header")
    expected_top = []
    for set_id, segment_set in segments["dimensionSets"].items():
        overall_by_id = {row["ipId"]: row for row in segment_set["overallIps"]}
        for sparse_group in segment_set["groups"]:
            group = materialize_group(segment_set, sparse_group)
            values = group["values"]
            prefix = [
                dimension_labels[set_id],
                values.get("age"),
                gender_labels.get(values.get("gender")),
                region_labels.get(values.get("region")),
                group["groupUsers"],
                group["eligibleUsers"],
                segment_set["overallEligibleUsers"],
            ]
            ranked = sorted(
                (row for row in group["ips"] if row["interestUsers"] > 0),
                key=lambda row: (-row["interestUsers"], -(row["groupCoverage"] or 0), row["ipId"]),
            )[:10]
            if not ranked:
                expected_top.append(prefix + [None, None, None, None, None, None, group["sampleStatus"]])
            for rank, row in enumerate(ranked, start=1):
                difference = row["percentagePointDifference"]
                expected_top.append(prefix + [
                    rank,
                    names[row["ipId"]],
                    row["interestUsers"],
                    row["groupCoverage"],
                    overall_by_id[row["ipId"]]["overallCoverage"],
                    difference * 100 if difference is not None else None,
                    group["sampleStatus"],
                ])
    for offset, row in enumerate(expected_top, start=3):
        _require_sheet_row(sheet, offset, row, "segment Top 10")
    _require(sheet.max_row == len(expected_top) + 2, "workbook segment Top 10 row count mismatch")


def verify(
    workbook_path: str | Path,
    detail_path: str | Path,
    snapshot_path: str | Path,
    candidate_path: str | Path,
    manifest_path: str | Path,
    registry_path: str | Path,
    private_detail_path: str | Path | None = None,
    base_detail_path: str | Path | None = None,
    segment_path: str | Path | None = None,
) -> VerificationResult:
    workbook_path = Path(workbook_path).resolve()
    detail_path = Path(detail_path).resolve()
    snapshot_path = Path(snapshot_path).resolve()
    candidate_path = Path(candidate_path).resolve()
    manifest_path = Path(manifest_path).resolve()
    registry_path = Path(registry_path).resolve()
    for path in (workbook_path, detail_path, snapshot_path, candidate_path, manifest_path, registry_path):
        _require(path.is_file(), f"missing verification input: {path.name}")

    manifest = _load_json(manifest_path, "interest manifest")
    manifest_schema = manifest.get("schemaVersion")
    _require(
        manifest_schema in {
            MANIFEST_SCHEMA,
            PREVIOUS_MANIFEST_SCHEMA,
            EARLIER_MANIFEST_SCHEMA,
            OLDER_MANIFEST_SCHEMA,
            LEGACY_MANIFEST_SCHEMA,
            INITIAL_MANIFEST_SCHEMA,
        },
        "interest manifest schema mismatch",
    )
    current_detail_contract = manifest_schema in {
        MANIFEST_SCHEMA,
        PREVIOUS_MANIFEST_SCHEMA,
        EARLIER_MANIFEST_SCHEMA,
        OLDER_MANIFEST_SCHEMA,
    }
    detail_schema = DETAIL_SCHEMA if current_detail_contract else LEGACY_DETAIL_SCHEMA
    detail_fields = DETAIL_FIELDS if current_detail_contract else LEGACY_DETAIL_FIELDS
    if manifest_schema in {MANIFEST_SCHEMA, PREVIOUS_MANIFEST_SCHEMA, EARLIER_MANIFEST_SCHEMA}:
        private_detail_schema = PRIVATE_DETAIL_SCHEMA
        private_detail_fields = PRIVATE_DETAIL_FIELDS
        default_private_name = "teeni-interest-detail.csv"
    elif manifest_schema == OLDER_MANIFEST_SCHEMA:
        private_detail_schema = PREVIOUS_PRIVATE_DETAIL_SCHEMA
        private_detail_fields = PREVIOUS_PRIVATE_DETAIL_FIELDS
        default_private_name = "teeni-interest-detail.private.csv"
    else:
        private_detail_schema = LEGACY_PRIVATE_DETAIL_SCHEMA
        private_detail_fields = LEGACY_PRIVATE_DETAIL_FIELDS
        default_private_name = "teeni-interest-detail.private.csv"
    resolved_private_detail: Path | None = None
    if manifest_schema in {
        MANIFEST_SCHEMA,
        PREVIOUS_MANIFEST_SCHEMA,
        EARLIER_MANIFEST_SCHEMA,
        OLDER_MANIFEST_SCHEMA,
        LEGACY_MANIFEST_SCHEMA,
    }:
        resolved_private_detail = Path(
            private_detail_path or manifest_path.parent / default_private_name
        ).resolve()
        _require(resolved_private_detail.is_file(), f"missing verification input: {resolved_private_detail.name}")
        _require(
            manifest.get("privateDetailSchema") == private_detail_schema,
            "private detail schema mismatch",
        )
    resolved_base_detail: Path | None = None
    if manifest_schema in {MANIFEST_SCHEMA, PREVIOUS_MANIFEST_SCHEMA}:
        resolved_base_detail = Path(base_detail_path).resolve() if base_detail_path else None
        _require(
            resolved_base_detail is not None and resolved_base_detail.is_file(),
            "base detail is required for current demographic interest manifests",
        )
        _require(
            sha256_file(resolved_base_detail) == manifest.get("input", {}).get("detailSha256"),
            "base detail hash mismatch",
        )
    resolved_segments: Path | None = None
    if manifest_schema == MANIFEST_SCHEMA:
        resolved_segments = Path(
            segment_path or manifest_path.parent / "teeni-interest-segments.json"
        ).resolve()
        _require(resolved_segments.is_file(), f"missing verification input: {resolved_segments.name}")
    snapshot = _load_json(snapshot_path, "interest snapshot")
    candidates = _load_json(candidate_path, "candidate package")
    scene_id, product_version = product_identity(snapshot)
    _require(product_identity(manifest) == (scene_id, product_version), "manifest product identity mismatch")
    _require(product_identity(candidates) == (scene_id, product_version), "candidate product identity mismatch")
    segments = _load_json(resolved_segments, "interest segments") if resolved_segments else None
    registry = EntityRegistry.from_json(registry_path)
    province_map = ProvinceMap.load(DEFAULT_PROVINCE_MAP)
    _require(manifest.get("engineVersion") == ENGINE_VERSION, "interest engine version mismatch")
    _require(snapshot.get("schemaVersion") == SNAPSHOT_SCHEMA, "interest snapshot schema mismatch")
    _require(snapshot.get("reportSchema") == REPORT_SCHEMA, "interest report schema mismatch")
    _require(snapshot.get("detailSchema") == detail_schema, "interest detail schema mismatch")
    _require(candidates.get("schemaVersion") == CANDIDATE_SCHEMA, "candidate package schema mismatch")
    _require(manifest.get("registry", {}).get("sha256") == sha256_file(registry_path), "registry hash mismatch")
    _require(snapshot.get("registrySha256") == sha256_file(registry_path), "snapshot registry hash mismatch")
    outputs = {item.get("role"): item for item in manifest.get("outputs", []) if isinstance(item, dict)}
    expected_outputs = {
        "workbook": workbook_path,
        "detail": detail_path,
        "snapshot": snapshot_path,
        "candidates_private": candidate_path,
    }
    if resolved_private_detail is not None:
        expected_outputs["private_detail"] = resolved_private_detail
    if resolved_segments is not None:
        expected_outputs["segments"] = resolved_segments
    _require(set(outputs) == set(expected_outputs), "interest manifest output roles mismatch")
    for role, path in expected_outputs.items():
        item = outputs[role]
        _require(item.get("file") == path.name, f"{role} filename mismatch")
        _require(item.get("sha256") == sha256_file(path), f"{role} hash mismatch")
        _require(item.get("bytes") == path.stat().st_size, f"{role} byte size mismatch")

    _assert_public_keys(snapshot)
    if segments is not None:
        _assert_public_keys(segments)
    _require(all("examples" not in row for row in snapshot.get("candidates", [])), "public candidates contain examples")
    private_rows = candidates.get("candidates")
    _require(isinstance(private_rows, list) and len(private_rows) <= 200, "candidate review package exceeds 200-item cap")
    public_rows = snapshot.get("candidates")
    _require(
        public_rows == [
            {key: value for key, value in row.items() if key != "examples"}
            for row in private_rows[:10]
        ],
        "public and private candidate aggregates do not reconcile",
    )
    min_users = int(snapshot.get("method", {}).get("candidateMinUsers", 5))
    min_sessions = int(snapshot.get("method", {}).get("candidateMinSessions", 8))
    for row in private_rows:
        _require(len(row.get("examples", [])) <= 3, "candidate evidence exceeds three examples")
        _require(not CONTACT_LIKE.search(str(row.get("phrase", ""))), "candidate phrase contains contact-like data")
        _require(str(row.get("suggestedType")) in {item.value for item in EntityType}, "candidate type is invalid")
        growth = row.get("sevenDayGrowth")
        volume_ok = int(row.get("users", 0)) >= min_users and int(row.get("sessions", 0)) >= min_sessions
        growth_ok = int(row.get("users", 0)) >= 3 and isinstance(growth, (int, float)) and growth >= 3
        _require(
            bool(row.get("dailyEligible")) == bool(volume_ok or growth_ok),
            "candidate daily eligibility flag is invalid",
        )

    recomputed = _recompute_detail(detail_path, registry, detail_fields)
    recomputed_segments = None
    if resolved_base_detail is not None:
        recomputed_segments = _recompute_segments(
            resolved_base_detail,
            detail_path,
            registry,
            str(manifest.get("input", {}).get("detailSha256") or ""),
            str(snapshot.get("dataDate") or ""),
            sha256_file(registry_path),
            province_map,
            scene_id,
        )
    private_detail_rows = 0
    if resolved_private_detail is not None:
        private_detail_rows = _verify_private_detail(
            resolved_private_detail,
            detail_path,
            private_detail_fields,
            detail_fields,
        )
        _require(private_detail_rows == recomputed["rows"], "private detail row total mismatch")
    totals = snapshot.get("totals", {})
    _require(totals.get("inputQueries") == recomputed["rows"], "input query total mismatch")
    _require(totals.get("validContentQueries") == recomputed["valid"], "valid content total mismatch")
    _require(totals.get("entityQueries") == recomputed["entityQueries"], "entity query total mismatch")
    _require(totals.get("activeInterestUsers") == recomputed["activeUsers"], "active user total mismatch")
    _require(manifest.get("input", {}).get("rows") == recomputed["rows"], "manifest row total mismatch")
    _require(snapshot.get("sourceSha256") == manifest.get("input", {}).get("detailSha256"), "source hash mismatch")
    if resolved_base_detail is not None:
        expected_demographic = legacy_demographic(recomputed_segments)
        _require(
            snapshot.get("demographicInterest") == expected_demographic,
            "demographic interest aggregates mismatch",
        )
    if segments is not None:
        rollup_ids = [row["id"] for row in snapshot.get("ipRollups", [])]
        validate_segment_payload(
            segments,
            data_date=str(snapshot.get("dataDate") or ""),
            source_sha256=str(snapshot.get("sourceSha256") or ""),
            registry_sha256=sha256_file(registry_path),
            province_map_sha256=province_map.sha256,
            rollup_ids=rollup_ids,
            product_version=product_version,
            scene_id=scene_id,
        )
        normalized_segments = {**segments, "productVersion": product_version, "sceneId": scene_id}
        _require(normalized_segments == recomputed_segments, "interest segment aggregates mismatch")
        _require(
            manifest.get("provinceMap") == {"version": province_map.version, "sha256": province_map.sha256},
            "province map manifest mismatch",
        )
        expected_segment_meta = {
            "schemaVersion": SEGMENT_SCHEMA,
            "provinceMapVersion": province_map.version,
            "provinceMapSha256": province_map.sha256,
            "sampleMinimumUsers": segments["sampleMinimumUsers"],
            "domains": segments["domains"],
            "dimensionSets": [
                {
                    "id": "_".join(dimensions),
                    "dimensions": list(dimensions),
                    "groupCount": len(segments["dimensionSets"]["_".join(dimensions)]["groups"]),
                }
                for dimensions in DIMENSION_SETS
            ],
        }
        _require(snapshot.get("segmentInterest") == expected_segment_meta, "snapshot segment metadata mismatch")
    for key, enum_type, counts in (
        ("routes", QueryRoute, recomputed["routes"]),
        ("topics", BroadTopic, recomputed["topics"]),
        ("behaviors", Behavior, recomputed["behaviors"]),
    ):
        rows = snapshot.get(key)
        _require(isinstance(rows, list) and len(rows) == len(enum_type), f"{key} coverage mismatch")
        _require({row["name"]: row["queries"] for row in rows} == {item.value: counts[item.value] for item in enum_type}, f"{key} counts mismatch")
    snapshot_sections = {}
    for section in ("entities", "restrictedEntities", "ipRollups"):
        actual_rows = {row["id"]: row for row in snapshot.get(section, [])}
        expected_rows = recomputed[section]
        _require(set(actual_rows) == set(expected_rows), f"{section} entity set mismatch")
        snapshot_sections[section] = actual_rows
        for entity_id, expected in expected_rows.items():
            row = actual_rows[entity_id]
            for field, value in expected.items():
                actual = row.get(field)
                if isinstance(value, float):
                    _require(isinstance(actual, (int, float)) and abs(actual - value) < 1e-12, f"{entity_id} {field} mismatch")
                else:
                    _require(actual == value, f"{entity_id} {field} mismatch")
    snapshot_structures = {row["name"]: row for row in snapshot.get("entityStructures", [])}
    _require(set(snapshot_structures) == set(recomputed["entityStructures"]), "entityStructures set mismatch")
    for name, expected in recomputed["entityStructures"].items():
        for field, value in expected.items():
            actual = snapshot_structures[name].get(field)
            if isinstance(value, float):
                _require(isinstance(actual, (int, float)) and abs(actual - value) < 1e-12, f"{name} {field} mismatch")
            else:
                _require(actual == value, f"{name} {field} mismatch")

    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        _require(tuple(workbook.sheetnames) == EXPECTED_SHEETS, "workbook sheet layout mismatch")
        if "productVersion" in snapshot:
            _require(workbook["兴趣结论总览"]["A2"].value == "产品版本"
                     and workbook["兴趣结论总览"]["B2"].value == product_version
                     and workbook["兴趣结论总览"]["C2"].value == f"scene {scene_id} / {snapshot['dataDate']}",
                     "workbook product identity mismatch")
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    _require(not (isinstance(cell.value, str) and cell.value.startswith("=")), "workbook contains formulas")
        _require(workbook["IP总览"].max_row == len(snapshot_sections["ipRollups"]) + 1, "workbook IP row count mismatch")
        if segments is not None:
            _verify_segment_sheet(workbook["画像分层与IP"], snapshot, segments)
        _require(workbook["实体明细"].max_row == len(snapshot_sections["entities"]) + 1, "workbook entity row count mismatch")
        _require(workbook["风险热梗"].max_row == len(snapshot_sections["restrictedEntities"]) + 1, "workbook risk row count mismatch")
        _require(workbook["新梗雷达"].max_row == len(public_rows) + 1, "workbook candidate row count mismatch")
    finally:
        workbook.close()
    return VerificationResult(
        True,
        recomputed["rows"],
        len(outputs),
        len(EXPECTED_SHEETS),
        private_detail_rows,
    )
