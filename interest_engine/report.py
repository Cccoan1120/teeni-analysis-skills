from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .segments import GENDERS, materialize_group


YELLOW = "F4C430"
INK = "25231F"
MUTED = "66635D"
SUBTLE = "F4F1E8"
WHITE = "FFFFFF"


def _sheet(
    workbook: Workbook,
    name: str,
    headers: list[str],
    rows: list[list],
    widths: list[int],
    number_formats: dict[int, str] | None = None,
) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(name="Microsoft YaHei UI", color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 28
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei UI", color=INK, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, number_format in (number_formats or {}).items():
        for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            for item in cell:
                if isinstance(item.value, (int, float)):
                    item.number_format = number_format
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    if rows:
        table = Table(displayName=f"T{len(workbook.worksheets):02d}", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)


def _segment_sheet(workbook: Workbook, snapshot: dict, segments: dict) -> None:
    sheet = workbook.create_sheet("画像分层与IP")
    names = {row["id"]: row["name"] for row in snapshot["ipRollups"]}
    gender_labels = dict(GENDERS)
    region_labels = {row["id"]: row["label"] for row in segments["domains"]["regions"]}
    dimension_labels = {
        "age": "年龄",
        "gender": "性别",
        "region": "地区",
        "age_gender": "年龄×性别",
        "age_region": "年龄×地区",
        "gender_region": "性别×地区",
        "age_gender_region": "年龄×性别×地区",
    }
    top_headers = [
        "维度口径", "年龄", "性别", "地区", "分组用户", "兴趣可分析用户", "当前口径总体可分析用户",
        "排名", "IP", "兴趣用户", "组内覆盖率", "当前口径总体覆盖率", "百分点差（pp）", "样本状态",
    ]
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(top_headers))
    sheet.cell(1, 1, "七种画像口径各分组 Top 10 IP；pp = 组内覆盖率 - 当前维度口径总体覆盖率")
    sheet.append(top_headers)
    for set_id, segment_set in segments["dimensionSets"].items():
        overall_by_id = {row["ipId"]: row for row in segment_set["overallIps"]}
        for sparse_group in segment_set["groups"]:
            group = materialize_group(segment_set, sparse_group)
            values = group["values"]
            prefix = [
                dimension_labels[set_id],
                values.get("age"),
                gender_labels.get(values.get("gender")),
                region_labels.get(values.get("region")),
                group["groupUsers"],
                group["eligibleUsers"],
                segment_set["overallEligibleUsers"],
            ]
            ranked = sorted(
                (row for row in group["ips"] if row["interestUsers"] > 0),
                key=lambda row: (-row["interestUsers"], -(row["groupCoverage"] or 0), row["ipId"]),
            )[:10]
            if not ranked:
                sheet.append(prefix + [None, None, None, None, None, None, group["sampleStatus"]])
                continue
            for rank, row in enumerate(ranked, start=1):
                overall = overall_by_id[row["ipId"]]
                difference = row["percentagePointDifference"]
                sheet.append(prefix + [
                    rank,
                    names[row["ipId"]],
                    row["interestUsers"],
                    row["groupCoverage"],
                    overall["overallCoverage"],
                    difference * 100 if difference is not None else None,
                    group["sampleStatus"],
                ])

    title = sheet.cell(1, 1)
    title.fill = PatternFill("solid", fgColor=YELLOW)
    title.font = Font(name="Microsoft YaHei UI", color=INK, bold=True, size=12)
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26
    for cell in sheet[2]:
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(name="Microsoft YaHei UI", color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 34
    for row in sheet.iter_rows(min_row=3):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei UI", color=INK, size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        if row[-1].value == "样本不足":
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=SUBTLE)
    for row in range(3, sheet.max_row + 1):
        for column in (11, 12):
            sheet.cell(row, column).number_format = "0.0%"
        sheet.cell(row, 13).number_format = '+0.0"pp";-0.0"pp";-'
    widths = [18, 8, 8, 14, 14, 18, 22, 8, 24, 14, 14, 22, 16, 12]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "I3"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def write_workbook(path: Path, snapshot: dict, registry, segments: dict) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    totals = snapshot["totals"]
    _sheet(workbook, "兴趣结论总览", ["指标", "值", "口径"], [
        ["产品版本", snapshot.get("productVersion", "M1"), f"scene {snapshot.get('sceneId', '488')} / {snapshot['dataDate']}"],
        ["输入query", totals["inputQueries"], "基础逐轮明细全部行"],
        ["有效内容query", totals["validContentQueries"], "清洗路由为有效内容"],
        ["已识别实体query", totals["entityQueries"], "至少命中一个已批准实体"],
        ["主动兴趣用户", totals["activeInterestUsers"], "至少一次主动发起或实质延续的当日独立用户"],
        ["画像分层口径", len(segments["dimensionSets"]), "年龄、性别、地区及其全部交叉组合分别使用各自稳定画像口径"],
        ["待审候选", len(snapshot["candidates"]), "不进入正式IP榜"],
        ["候选补充状态", "降级" if snapshot["method"]["candidateDegraded"] else "完成", snapshot["method"].get("candidateFailureCode") or "无"],
    ], [24, 18, 68])
    _sheet(workbook, "IP总览", [
        "排名", "IP", "层级", "子类型", "来源", "命中子实体数", "主动兴趣用户", "提及用户", "主动发起用户", "主动延续用户",
        "会话", "query", "三轮以上深聊率", "七日变化", "明确正向偏好用户", "明确负向偏好用户", "无明确偏好提及用户", "同轮混合偏好用户", "IP提及样本", "样本状态",
    ], [[
        index, row["name"], "父级汇总" if row["rollupKind"] == "parent" else "独立IP",
        row["entitySubtype"], row["sourceLabel"], row["childEntityCount"],
        row["activeInterestUsers"], row["mentionUsers"],
        row["initiatorUsers"], row["continuationUsers"], row["sessions"], row["queries"], row["deepChatRate"],
        row["sevenDayChange"] if row["sevenDayChange"] is not None else "无基线",
        row["positivePreferenceUsers"], row["negativePreferenceUsers"], row["neutralMentionUsers"], row["mixedPreferenceUsers"], row["sampleUsers"], row["sampleStatus"],
    ] for index, row in enumerate(snapshot["ipRollups"], start=1)], [8, 24, 14, 18, 22, 14, 16, 14, 16, 16, 12, 12, 16, 14, 18, 18, 20, 20, 14, 14], {13: "0.0%", 14: "+0.0%;-0.0%;-"})
    _segment_sheet(workbook, snapshot, segments)
    _sheet(workbook, "实体明细", [
        "排名", "实体", "父IP", "类型", "子类型", "来源", "主动兴趣用户", "提及用户", "主动发起用户",
        "主动延续用户", "会话", "query", "三轮以上深聊率", "七日变化",
    ], [[
        index, row["name"], row["parentName"] or "-", row["entityType"], row["entitySubtype"],
        row["sourceLabel"], row["activeInterestUsers"], row["mentionUsers"], row["initiatorUsers"],
        row["continuationUsers"], row["sessions"], row["queries"], row["deepChatRate"],
        row["sevenDayChange"] if row["sevenDayChange"] is not None else "无基线",
    ] for index, row in enumerate(snapshot["entities"], start=1)], [8, 24, 24, 14, 18, 22, 16, 14, 16, 16, 12, 12, 16, 14], {13: "0.0%", 14: "+0.0%;-0.0%;-"})
    _sheet(workbook, "风险热梗", [
        "实体", "子类型", "安全分类", "来源", "主动兴趣用户", "提及用户", "主动发起用户", "主动延续用户",
        "会话", "query", "三轮以上深聊率", "七日变化",
    ], [[
        row["name"], row["entitySubtype"], row["safetyCategory"], row["sourceLabel"], row["activeInterestUsers"],
        row["mentionUsers"], row["initiatorUsers"], row["continuationUsers"], row["sessions"], row["queries"],
        row["deepChatRate"], row["sevenDayChange"] if row["sevenDayChange"] is not None else "无基线",
    ] for row in snapshot["restrictedEntities"]], [26, 18, 20, 22, 16, 14, 16, 16, 12, 12, 16, 14], {11: "0.0%", 12: "+0.0%;-0.0%;-"})
    _sheet(workbook, "新梗雷达", [
        "候选", "建议规范名", "建议类型", "独立用户", "会话", "query", "七日增长", "模型置信度", "状态",
    ], [[
        row["phrase"], row["suggestedName"], row["suggestedType"], row["users"], row["sessions"], row["queryCount"],
        row["sevenDayGrowth"] if row["sevenDayGrowth"] is not None else "无基线",
        row["modelConfidence"], "待人工审核",
    ] for row in snapshot["candidates"]], [24, 24, 16, 12, 12, 12, 14, 14, 16], {7: "0.0x"})
    _sheet(workbook, "普通话题", ["内容大类", "query", "占有效内容比例"], [
        [row["name"], row["queries"], row["rate"]] for row in snapshot["topics"]
    ], [28, 16, 20], {3: "0.0%"})
    _sheet(workbook, "用户行为", ["用户行为", "query", "占有效内容比例"], [
        [row["name"], row["queries"], row["rate"]] for row in snapshot["behaviors"]
    ], [28, 16, 20], {3: "0.0%"})
    _sheet(workbook, "清洗与路由", ["路由结果", "query", "占全部输入比例", "是否进入内容分析"], [
        [row["name"], row["queries"], row["rate"], "是" if row["name"] == "有效内容" else "否"] for row in snapshot["routes"]
    ], [24, 16, 20, 22], {3: "0.0%"})
    _sheet(workbook, "实体词典", [
        "实体ID", "规范名", "父IP", "类型", "子类型", "内容大类", "来源", "来源平台", "可见性", "安全分类", "自动规则", "待确认规则",
    ], [[
        entity.id,
        entity.canonical_name,
        registry.get(entity.parent_registry_id).canonical_name if entity.parent_registry_id else "-",
        entity.entity_type.value,
        entity.entity_subtype.value,
        entity.broad_topic.value,
        entity.source_label,
        entity.source_platform,
        entity.visibility.value,
        entity.safety_category.value,
        "、".join(alias.value for alias in entity.aliases if not alias.ambiguous),
        "、".join(alias.value for alias in entity.aliases if alias.ambiguous),
    ] for entity in registry.entities], [14, 24, 24, 14, 18, 24, 22, 18, 14, 18, 42, 36])
    _sheet(workbook, "分析口径", ["项目", "定义"], [
        ["产品版本", f"{snapshot.get('productVersion', 'M1')} / scene {snapshot.get('sceneId', '488')}；各版本独立统计。"],
        ["主动兴趣用户", "用户当日主动引入实体，或对上下文实体进行实质提问、描述、选择和延续；AI带出后的单纯确认与复述不计入。"],
        ["偏好方向", "逐实体记录明确正向、负向、同轮混合与无明确偏好。关注不等于喜欢；同一用户跨轮可属于多类，不可相加作为用户总数。"],
        ["上下文归因", "继续讲、然后呢等仅在上一轮用户与AI合计只有一个明确实体时继承；多实体歧义保留上下文不足。"],
        ["三轮以上深聊率", "同一会话内至少三个有效内容query命中同一实体的会话数 / 该实体全部命中会话数。"],
        ["IP总览", "父级汇总按父IP及其命中子实体重新去重；没有子实体但类型明确的作品、角色、游戏、玩具品牌和创作者账号作为独立IP展示。"],
        ["画像分层", "提供年龄、性别、地区、年龄×性别、年龄×地区、性别×地区、年龄×性别×地区七种口径；每种口径只要求所选维度画像稳定。"],
        ["地区口径", f"城市按冻结映射 {segments['provinceMapVersion']} 汇总到34个省级地区；无法映射的正常城市只进入诊断，不进入地区分组。"],
        ["百分点差（pp）", "pp 表示百分点，不是百分比变化率。pp = 组内覆盖率 - 当前维度口径总体覆盖率，例如 3.4% - 2.7% = +0.7pp。"],
        ["样本边界", "IP榜按该IP提及用户是否达到30标记样本；画像覆盖率分母为相应分组可分析用户，深聊率分母为该IP命中会话。零分母显示为空值。"],
        ["候选雷达", "未批准的新词只用于审核，不进入正式IP热度榜。"],
        ["准确性边界", "自动命中率与模型建议表示覆盖，不代表人工确认准确率。"],
        ["隐私边界", "工作簿和聚合快照不包含原始query、AI回复、clientId或cid。"],
    ], [24, 92])
    workbook.save(path)
