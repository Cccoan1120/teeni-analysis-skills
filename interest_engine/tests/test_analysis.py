import csv
import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from interest_engine.analysis import age_for_profile, analyze
from interest_engine.contracts import InputContractError, load_and_verify_input, sha256_file
from interest_engine.verify import verify


DETAIL_FIELDS = [
    "source_row", "clientId", "cid", "sceneId", "text", "ai_text", "turn_index",
    "is_template", "is_invalid_turn", "invalid_reason", "profile_age", "profile_gender",
    "age_status", "gender_status", "city_normalized", "city_status",
]


class BrokenEnricher:
    def enrich(self, _candidates):
        raise TimeoutError("offline")


class InterestAnalysisTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.detail = self.root / "base-detail.csv"
        self.base_manifest = self.root / "base-manifest.json"
        self.registry = self.root / "registry.json"
        self.registry.write_text(json.dumps({
            "schemaVersion": "teeni-interest-registry/1.0.0",
            "registryVersion": "test-1",
            "entities": [{
                "id": "IE0001",
                "canonicalName": "奥特曼",
                "entityType": "作品",
                "broadTopic": "影视动漫与角色",
                "aliases": ["Ultraman"],
                "ambiguousAliases": [],
            }],
        }, ensure_ascii=False), encoding="utf-8")
        rows = [
            self.row(1, "u1", "s1", "奥特曼", "好的"),
            self.row(2, "u1", "s1", "奥特曼是谁", "他是英雄"),
            self.row(3, "u1", "s1", "奥特曼为什么厉害", "因为他很勇敢"),
            self.row(4, "u2", "s2", "给我讲一个故事", "你喜欢奥特曼吗"),
            self.row(5, "u2", "s2", "奥特曼", "好"),
            self.row(6, "u3", "s3", "野生狗奶", ""),
            self.row(7, "u4", "s4", "野生狗奶", ""),
            self.row(8, "u5", "s5", "野生狗奶", ""),
            self.row(9, "u6", "s6", "嗯", ""),
            self.row(10, "u7", "s7", "调大音量", ""),
            self.row(11, "u8", "s8", "啊啊啊啊", ""),
            self.row(12, "u9", "s9", "继续", ""),
            self.row(13, "u10", "s10", "奥特曼", "", is_template="是"),
        ]
        with self.detail.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=DETAIL_FIELDS)
            writer.writeheader()
            writer.writerows(rows)
        detail_hash = hashlib.sha256(self.detail.read_bytes()).hexdigest()
        self.base_manifest.write_text(json.dumps({
            "schemaVersion": "teeni-base-bundle-manifest/1.0.0",
            "contractVersion": "teeni-base-detail/1.2.0",
            "sceneId": "488",
            "outputs": [{
                "role": "primary_detail",
                "file": self.detail.name,
                "sha256": detail_hash,
                "rows": len(rows),
            }],
        }), encoding="utf-8")

    @staticmethod
    def row(
        index,
        user,
        session,
        text,
        ai_text,
        *,
        is_template="否",
        age="6",
        gender="男",
        age_status="正常",
        gender_status="正常",
        city="北京",
        city_status="正常",
    ):
        return {
            "source_row": index,
            "clientId": user,
            "cid": session,
            "sceneId": "488",
            "text": text,
            "ai_text": ai_text,
            "turn_index": index,
            "is_template": is_template,
            "is_invalid_turn": "否",
            "invalid_reason": "",
            "profile_age": age,
            "profile_gender": gender,
            "age_status": age_status,
            "gender_status": gender_status,
            "city_normalized": city,
            "city_status": city_status,
        }

    def write_rows(self, rows):
        with self.detail.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=DETAIL_FIELDS)
            writer.writeheader()
            writer.writerows(rows)
        payload = json.loads(self.base_manifest.read_text(encoding="utf-8"))
        payload["outputs"][0]["rows"] = len(rows)
        payload["outputs"][0]["sha256"] = hashlib.sha256(self.detail.read_bytes()).hexdigest()
        self.base_manifest.write_text(json.dumps(payload), encoding="utf-8")

    def run_analysis(self, name="output", enricher=None):
        return analyze(
            self.detail,
            self.base_manifest,
            self.registry,
            self.root / name,
            data_date="2026-08-25",
            enricher=enricher,
            candidate_min_users=3,
            candidate_min_sessions=3,
        )

    def test_analysis_metrics_candidates_and_independent_verification(self):
        artifacts = self.run_analysis()
        snapshot = artifacts.snapshot
        self.assertEqual(snapshot["totals"], {
            "inputQueries": 13,
            "validContentQueries": 8,
            "entityQueries": 4,
            "activeInterestUsers": 1,
        })
        entity = snapshot["entities"][0]
        self.assertEqual(entity["name"], "奥特曼")
        self.assertEqual(entity["activeInterestUsers"], 1)
        self.assertEqual(entity["mentionUsers"], 2)
        self.assertEqual(entity["initiatorUsers"], 1)
        self.assertEqual(entity["continuationUsers"], 1)
        self.assertEqual(entity["queries"], 4)
        self.assertEqual(entity["deepSessions"], 1)
        self.assertEqual(entity["deepChatRate"], 0.5)
        self.assertEqual(snapshot["candidates"][0]["phrase"], "野生狗奶")
        self.assertNotIn("examples", snapshot["candidates"][0])
        self.assertTrue(snapshot["method"]["candidateDegraded"])

        self.assertFalse((artifacts.snapshot_path.parent / ".interest-work.private.sqlite3").exists())
        with artifacts.private_detail_path.open("r", encoding="utf-8-sig", newline="") as stream:
            private_rows = list(csv.DictReader(stream))
        self.assertEqual(len(private_rows), 13)
        self.assertEqual(private_rows[0]["clientId"], "u1")
        self.assertEqual(private_rows[0]["cid"], "s1")
        self.assertEqual(private_rows[0]["text"], "奥特曼")
        self.assertEqual(private_rows[0]["ai_text"], "好的")
        self.assertEqual(private_rows[0]["entity_names"], "奥特曼")
        self.assertNotIn("entity_ids", private_rows[0])
        self.assertNotIn("user_key", private_rows[0])
        with artifacts.detail_path.open("r", encoding="utf-8-sig", newline="") as stream:
            safe_rows = list(csv.DictReader(stream))
        self.assertEqual(safe_rows[0]["entity_names"], "奥特曼")
        self.assertEqual(safe_rows[0]["entity_ids"], "IE0001")
        self.assertEqual(artifacts.private_detail_path.name, "teeni-interest-detail.csv")
        self.assertEqual(artifacts.detail_path.name, "teeni-interest-detail.safe.csv")
        self.assertEqual(artifacts.segment_path.name, "teeni-interest-segments.json")
        segments = json.loads(artifacts.segment_path.read_text(encoding="utf-8"))
        self.assertEqual(segments["schemaVersion"], "teeni-interest-segments/1.0.0")
        self.assertEqual(
            {key: len(value["groups"]) for key, value in segments["dimensionSets"].items()},
            {
                "age": 17,
                "gender": 2,
                "region": 34,
                "age_gender": 34,
                "age_region": 578,
                "gender_region": 68,
                "age_gender_region": 1156,
            },
        )

        result = verify(
            artifacts.workbook_path,
            artifacts.detail_path,
            artifacts.snapshot_path,
            artifacts.candidate_path,
            artifacts.manifest_path,
            self.registry,
            base_detail_path=self.detail,
        )
        self.assertTrue(result.ok)
        self.assertEqual(result.detail_rows, 13)
        self.assertEqual(result.private_detail_rows, 13)
        self.assertEqual(result.sheet_count, 11)

        workbook = load_workbook(artifacts.workbook_path, read_only=False, data_only=False)
        try:
            self.assertEqual(workbook["实体明细"]["M2"].number_format, "0.0%")
            self.assertEqual(workbook["普通话题"]["C2"].number_format, "0.0%")
            self.assertEqual(workbook["清洗与路由"]["C2"].number_format, "0.0%")
            self.assertEqual(workbook["实体明细"]["N2"].value, "无基线")
            self.assertEqual(workbook["新梗雷达"]["G2"].value, "无基线")
        finally:
            workbook.close()

    def test_cooccurring_unknown_candidate_and_entity_preferences(self):
        rows = [self.row(index, f"user-{index}", f"session-{index}", "我喜欢奥特曼，也喜欢《跨域新作品》", "好的")
                for index in range(1, 4)]
        self.write_rows(rows)
        artifacts = self.run_analysis("cooccurring")
        self.assertIn("跨域新作品", [row["phrase"] for row in artifacts.snapshot["candidates"]])
        ip = artifacts.snapshot["ipRollups"][0]
        self.assertEqual(ip["positivePreferenceUsers"], 3)
        self.assertEqual(ip["sampleUsers"], 3)
        self.assertEqual(ip["sampleStatus"], "样本不足")

    def test_candidate_model_failure_degrades_without_failing_analysis(self):
        artifacts = self.run_analysis("degraded", enricher=BrokenEnricher())
        self.assertTrue(artifacts.snapshot["method"]["candidateDegraded"])
        self.assertEqual(artifacts.snapshot["method"]["candidateFailureCode"], "candidate_model_unavailable")
        self.assertEqual(artifacts.snapshot["candidates"][0]["modelConfidence"], "失败")

    def test_manifest_hash_mismatch_is_rejected_before_analysis(self):
        payload = json.loads(self.base_manifest.read_text(encoding="utf-8"))
        payload["outputs"][0]["sha256"] = "0" * 64
        self.base_manifest.write_text(json.dumps(payload), encoding="utf-8")
        with self.assertRaises(InputContractError):
            self.run_analysis("rejected")

    def test_m2_detects_source_and_verifies_independently(self):
        rows = [self.row(1, "u1", "s1", "奥特曼", "")]
        rows[0]["sceneId"] = "904"
        self.write_rows(rows)
        payload = json.loads(self.base_manifest.read_text(encoding="utf-8"))
        payload["sceneId"] = "904"
        self.base_manifest.write_text(json.dumps(payload), encoding="utf-8")
        contract = load_and_verify_input(self.detail, self.base_manifest)
        self.assertEqual((contract.scene_id, contract.product_version), ("904", "M2"))
        artifacts = self.run_analysis("m2")
        self.assertEqual((artifacts.snapshot["sceneId"], artifacts.snapshot["productVersion"]), ("904", "M2"))
        segments = json.loads(artifacts.segment_path.read_text(encoding="utf-8"))
        self.assertEqual((segments["sceneId"], segments["productVersion"]), ("904", "M2"))
        self.assertTrue(verify(artifacts.workbook_path, artifacts.detail_path, artifacts.snapshot_path,
                               artifacts.candidate_path, artifacts.manifest_path, self.registry,
                               base_detail_path=self.detail).ok)
        with self.assertRaisesRegex(InputContractError, "same product"):
            analyze(self.detail, self.base_manifest, self.registry, self.root / "mixed-history",
                    data_date="2026-08-25", history=[{"dataDate": "2026-08-24"}])
        segments.update(sceneId="488", productVersion="M1")
        artifacts.segment_path.write_text(json.dumps(segments), encoding="utf-8")
        manifest = json.loads(artifacts.manifest_path.read_text(encoding="utf-8"))
        output = next(row for row in manifest["outputs"] if row["role"] == "segments")
        output.update(sha256=sha256_file(artifacts.segment_path), bytes=artifacts.segment_path.stat().st_size)
        artifacts.manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "segment product version mismatch"):
            verify(artifacts.workbook_path, artifacts.detail_path, artifacts.snapshot_path,
                   artifacts.candidate_path, artifacts.manifest_path, self.registry, base_detail_path=self.detail)

    def test_scene_row_mismatch_and_unsupported_scene_are_rejected(self):
        rows = [self.row(1, "u1", "s1", "奥特曼", ""), self.row(2, "u2", "s2", "奥特曼", "")]
        rows[1]["sceneId"] = "904"
        self.write_rows(rows)
        with self.assertRaisesRegex(InputContractError, "mixed scenes"):
            self.run_analysis("mixed")
        payload = json.loads(self.base_manifest.read_text(encoding="utf-8"))
        payload["sceneId"] = "901"
        self.base_manifest.write_text(json.dumps(payload), encoding="utf-8")
        with self.assertRaisesRegex(InputContractError, "488.*904"):
            load_and_verify_input(self.detail, self.base_manifest)

    def test_legacy_m1_outputs_without_identity_still_verify(self):
        artifacts = self.run_analysis("legacy-m1")
        for path in (artifacts.snapshot_path, artifacts.candidate_path, artifacts.manifest_path, artifacts.segment_path):
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload.pop("productVersion")
            payload.pop("sceneId")
            path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        manifest = json.loads(artifacts.manifest_path.read_text(encoding="utf-8"))
        for output in manifest["outputs"]:
            path = artifacts.manifest_path.parent / output["file"]
            output.update(sha256=sha256_file(path), bytes=path.stat().st_size)
        artifacts.manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        self.assertTrue(verify(artifacts.workbook_path, artifacts.detail_path, artifacts.snapshot_path,
                               artifacts.candidate_path, artifacts.manifest_path, self.registry,
                               base_detail_path=self.detail).ok)

    def test_parent_rollup_deduplicates_queries_and_restricted_entities_are_isolated(self):
        self.registry.write_text(json.dumps({
            "schemaVersion": "teeni-interest-registry/1.1.0",
            "registryVersion": "hierarchy-test",
            "entities": [
                {
                    "id": "IP1", "canonicalName": "奥特曼", "entityType": "作品",
                    "entitySubtype": "特摄", "broadTopic": "影视动漫与角色",
                    "matchRules": [{"value": "奥特曼", "mode": "substring", "policy": "auto"}],
                },
                {
                    "id": "C1", "canonicalName": "赛罗", "entityType": "角色",
                    "entitySubtype": "作品角色", "broadTopic": "影视动漫与角色",
                    "parentRegistryId": "IP1",
                    "matchRules": [{"value": "赛罗", "mode": "substring", "policy": "auto"}],
                },
                {
                    "id": "C2", "canonicalName": "迪迦", "entityType": "角色",
                    "entitySubtype": "作品角色", "broadTopic": "影视动漫与角色",
                    "parentRegistryId": "IP1",
                    "matchRules": [{"value": "迪迦", "mode": "substring", "policy": "auto"}],
                },
                {
                    "id": "R1", "canonicalName": "风险固定表达", "entityType": "网络热梗",
                    "entitySubtype": "风险事件", "broadTopic": "其他明确内容",
                    "visibility": "restricted", "safetyCategory": "real_case",
                    "matchRules": [{"value": "风险固定表达", "mode": "substring", "policy": "auto"}],
                },
            ],
        }, ensure_ascii=False), encoding="utf-8")
        rows = [
            self.row(1, "u1", "s1", "赛罗和迪迦", ""),
            self.row(2, "u1", "s1", "赛罗", ""),
            self.row(3, "u1", "s1", "迪迦", ""),
            self.row(4, "u2", "s2", "风险固定表达", ""),
        ]
        self.write_rows(rows)

        snapshot = self.run_analysis("hierarchy").snapshot

        self.assertEqual({row["id"] for row in snapshot["entities"]}, {"C1", "C2"})
        self.assertEqual({row["id"] for row in snapshot["restrictedEntities"]}, {"R1"})
        rollup = snapshot["ipRollups"][0]
        self.assertEqual(rollup["id"], "IP1")
        self.assertEqual(rollup["mentionUsers"], 1)
        self.assertEqual(rollup["sessions"], 1)
        self.assertEqual(rollup["queries"], 3)
        self.assertEqual(rollup["deepSessions"], 1)
        self.assertEqual(rollup["childEntityCount"], 2)
        self.assertEqual(rollup["rollupKind"], "parent")
        self.assertEqual(snapshot["totals"]["entityQueries"], 3)
        self.assertEqual(snapshot["totals"]["activeInterestUsers"], 1)

    def test_ip_rollups_include_only_explicit_standalone_ip_types(self):
        self.registry.write_text(json.dumps({
            "schemaVersion": "teeni-interest-registry/1.1.0",
            "registryVersion": "standalone-ip-test",
            "entities": [
                {
                    "id": "IP1", "canonicalName": "奥特曼", "entityType": "作品",
                    "entitySubtype": "特摄", "broadTopic": "影视动漫与角色",
                    "matchRules": [{"value": "奥特曼", "mode": "substring", "policy": "auto"}],
                },
                {
                    "id": "C1", "canonicalName": "赛罗", "entityType": "角色",
                    "entitySubtype": "作品角色", "broadTopic": "影视动漫与角色",
                    "parentRegistryId": "IP1",
                    "matchRules": [{"value": "赛罗", "mode": "substring", "policy": "auto"}],
                },
                {
                    "id": "IP2", "canonicalName": "汪汪队立大功", "entityType": "作品",
                    "entitySubtype": "动画", "broadTopic": "影视动漫与角色",
                    "matchRules": [{"value": "汪汪队", "mode": "substring", "policy": "auto"}],
                },
                {
                    "id": "IP3", "canonicalName": "哪吒", "entityType": "角色",
                    "entitySubtype": "作品角色", "broadTopic": "影视动漫与角色",
                    "matchRules": [{"value": "哪吒", "mode": "substring", "policy": "auto"}],
                },
                {
                    "id": "A1", "canonicalName": "猜谜语", "entityType": "游戏",
                    "entitySubtype": "其他热梗", "broadTopic": "电子游戏与互动玩法",
                    "matchRules": [{"value": "猜谜语", "mode": "substring", "policy": "auto"}],
                },
            ],
        }, ensure_ascii=False), encoding="utf-8")
        rows = [
            self.row(1, "u1", "s1", "赛罗", ""),
            self.row(2, "u2", "s2", "汪汪队", ""),
            self.row(3, "u3", "s3", "哪吒", ""),
            self.row(4, "u4", "s4", "猜谜语", ""),
        ]
        self.write_rows(rows)

        snapshot = self.run_analysis("standalone-ip").snapshot

        self.assertEqual(
            {row["id"]: row["rollupKind"] for row in snapshot["ipRollups"]},
            {"IP1": "parent", "IP2": "standalone", "IP3": "standalone"},
        )
        self.assertEqual(next(row for row in snapshot["ipRollups"] if row["id"] == "IP1")["childEntityCount"], 1)
        self.assertNotIn("A1", {row["id"] for row in snapshot["ipRollups"]})

    def test_exact_age_boundaries(self):
        self.assertEqual(
            [age_for_profile(value) for value in range(1, 18)],
            list(range(1, 18)),
        )
        for value in (None, "", 0, 18, "6.5", "not-an-age"):
            self.assertIsNone(age_for_profile(value))

    def test_demographic_interest_filters_profiles_and_deduplicates_users_and_parent_ip(self):
        self.registry.write_text(json.dumps({
            "schemaVersion": "teeni-interest-registry/1.1.0",
            "registryVersion": "demographic-test",
            "entities": [
                {
                    "id": "IP1", "canonicalName": "奥特曼", "entityType": "作品",
                    "entitySubtype": "特摄", "broadTopic": "影视动漫与角色",
                    "matchRules": [{"value": "奥特曼", "mode": "substring", "policy": "auto"}],
                },
                {
                    "id": "C1", "canonicalName": "赛罗", "entityType": "角色",
                    "entitySubtype": "作品角色", "broadTopic": "影视动漫与角色",
                    "parentRegistryId": "IP1",
                    "matchRules": [{"value": "赛罗", "mode": "substring", "policy": "auto"}],
                },
                {
                    "id": "C2", "canonicalName": "迪迦", "entityType": "角色",
                    "entitySubtype": "作品角色", "broadTopic": "影视动漫与角色",
                    "parentRegistryId": "IP1",
                    "matchRules": [{"value": "迪迦", "mode": "substring", "policy": "auto"}],
                },
            ],
        }, ensure_ascii=False), encoding="utf-8")
        rows = []
        source_row = 1
        for index in range(28):
            rows.append(self.row(source_row, f"m{index}", f"ms{index}", "给我讲一个故事", "", age="3", gender="男"))
            source_row += 1
        rows.append(self.row(source_row, "m28", "ms28", "赛罗和迪迦", "", age="3", gender="男"))
        source_row += 1
        rows.append(self.row(source_row, "m29", "ms29", "赛罗", "", age="3", gender="男"))
        source_row += 1
        rows.append(self.row(source_row, "m28", "ms28-second", "迪迦", "", age="3", gender="男"))
        source_row += 1
        rows.append(self.row(source_row, "f1", "fs1", "赛罗", "", age="2", gender="女"))
        source_row += 1
        rows.append(self.row(source_row, "f1", "fs2", "迪迦", "", age="2", gender="女"))
        source_row += 1
        for index in range(30):
            rows.append(self.row(
                source_row,
                f"older-f{index}",
                f"older-fs{index}",
                "奥特曼",
                "",
                age="10",
                gender="女",
                is_template="是",
            ))
            source_row += 1
        rows.extend([
            self.row(
                source_row, "missing-age", "x1", "赛罗", "", age="", gender="男",
                age_status="缺失", city="不存在的城市",
            ),
            self.row(source_row + 1, "conflict-gender", "x2", "赛罗", "", age="4", gender="", gender_status="冲突"),
            self.row(source_row + 2, "adult", "x3", "赛罗", "", age="18", gender="女"),
            self.row(source_row + 3, "unstable-profile", "x4", "赛罗", "", age="6", gender="男"),
            self.row(source_row + 4, "unstable-profile", "x5", "赛罗", "", age="7", gender="男"),
        ])
        self.write_rows(rows)

        artifacts = self.run_analysis("demographic")
        snapshot = artifacts.snapshot
        demographic = snapshot["demographicInterest"]
        self.assertEqual(demographic["overallEligibleUsers"], 31)
        self.assertEqual([row["ipId"] for row in demographic["overallIps"]], ["IP1"])
        self.assertEqual(demographic["overallIps"][0]["interestUsers"], 3)
        self.assertAlmostEqual(demographic["overallIps"][0]["overallCoverage"], 3 / 31)
        self.assertEqual(len(demographic["groups"]), 34)
        self.assertEqual(
            [(row["age"], row["gender"]) for row in demographic["groups"]],
            [(age, gender) for age in range(1, 18) for gender in ("男", "女")],
        )

        groups = {(row["age"], row["gender"]): row for row in demographic["groups"]}
        male = groups[(3, "男")]
        self.assertEqual((male["groupUsers"], male["eligibleUsers"], male["sampleStatus"]), (30, 30, "可描述"))
        self.assertEqual(male["ips"][0]["interestUsers"], 2)
        self.assertAlmostEqual(male["ips"][0]["groupCoverage"], 2 / 30)
        self.assertAlmostEqual(male["ips"][0]["percentagePointDifference"], 2 / 30 - 3 / 31)

        young_female = groups[(2, "女")]
        self.assertEqual((young_female["groupUsers"], young_female["eligibleUsers"], young_female["sampleStatus"]), (1, 1, "样本不足"))
        self.assertEqual(young_female["ips"][0]["interestUsers"], 1)

        zero_denominator = groups[(10, "女")]
        self.assertEqual((zero_denominator["groupUsers"], zero_denominator["eligibleUsers"], zero_denominator["sampleStatus"]), (30, 0, "样本不足"))
        self.assertIsNone(zero_denominator["ips"][0]["groupCoverage"])
        self.assertIsNone(zero_denominator["ips"][0]["percentagePointDifference"])

        segments = json.loads(artifacts.segment_path.read_text(encoding="utf-8"))
        self.assertEqual(segments["mappingDiagnostics"]["unmappedUsers"], 1)
        self.assertGreater(
            segments["dimensionSets"]["gender"]["overallEligibleUsers"],
            segments["dimensionSets"]["age_gender"]["overallEligibleUsers"],
        )
        male_summary = next(
            group for group in segments["dimensionSets"]["gender"]["groups"]
            if group["groupKey"] == "male"
        )
        self.assertGreater(male_summary["eligibleUsers"], male["eligibleUsers"])

        workbook = load_workbook(artifacts.workbook_path, read_only=True, data_only=True)
        try:
            rows = list(workbook["画像分层与IP"].iter_rows(min_row=3, values_only=True))
            sample_row = next(row for row in rows if row[0] == "年龄×性别" and row[1] == 2 and row[2] == "女")
            self.assertEqual((sample_row[7], sample_row[8], sample_row[13]), (1, "奥特曼", "样本不足"))
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
