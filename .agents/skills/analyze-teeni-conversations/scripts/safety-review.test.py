import csv
import importlib.util
import json
import tempfile
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook

spec = importlib.util.spec_from_file_location("safety_review", Path(__file__).with_name("safety-review.py"))
safety = importlib.util.module_from_spec(spec)
spec.loader.exec_module(safety)


def fixture():
    detail = [
        {"id": str(index), "clientId": "00001", "cid": "s1", "source_row": str(index + 1),
         "turn_index": str(index), "created_at": "2026-09-08 12:00:00", "timestamp": "1788840000",
         "text": text, "ai_text": reply, "safety_signals": "风险(高)" if index == 2 else "",
         "user_safety_signals": "用户风险(高)" if index == 2 else ""}
        for index, text, reply in [(1, "玩数字炸弹", "来玩数字炸弹"), (2, '=HYPERLINK("bad","bad")', "普通回复"), (3, "后续内容", "后续回答")]
    ]
    base = {"record_id": "2", "client_id": "00001", "cid": "s1", "source_time": "2026-09-08 12:00:00", "category": "风险", "severity": "高", "evidence": "词", "reason": "候选依据"}
    report = {"coreVersion": "2.6.1", "rulesVersion": "15.0.0", "primary": {
        "sceneId": "488", "summary": {"rows": 3}, "safetyIssues": [base, {**base, "category": "另一类别"}],
        "userSafetyIssues": [{**base, "category": "用户风险", "reviewPriority": "P2"}],
        "attentionUsers": [{"clientId": "00001", "reviewPriority": "P2", "priorityReason": "一次命中"}],
        "safetyGameExclusions": [{"record_id": "1", "client_id": "00001", "cid": "s1", "side": "user", "start": 3, "end": 5, "matched_token": "炸弹", "context_evidence": "玩数字炸弹", "reason": "游戏"}],
    }}
    return report, detail


class SafetyWorkbookTest(unittest.TestCase):
    def build(self, report, detail, folder):
        output, evidence = Path(folder) / "review.xlsx", Path(folder) / "exclusions.json"
        result = safety.build(report, detail, "2026-09-08", str(output), str(evidence))
        self.assertTrue(result["verified"])
        return output, evidence

    def test_multicategory_complete_context_and_literal_formulas(self):
        report, detail = fixture()
        with tempfile.TemporaryDirectory() as folder:
            output, evidence = self.build(report, detail, folder)
            workbook = load_workbook(output)
            self.assertEqual(workbook["AI回复候选"].max_row, 3)
            self.assertEqual(workbook["会话上下文"].max_row, 4)
            self.assertEqual(workbook["AI回复候选"]["J2"].data_type, "s")
            self.assertEqual(workbook["AI回复候选"]["A2"].value, "00001")
            self.assertEqual(workbook["AI回复候选"]["N2"].value, "待复核")
            self.assertIsNone(workbook["AI回复候选"]["O2"].value)
            self.assertEqual(workbook["AI回复候选"]["Q2"].hyperlink.location, "'会话上下文'!A3")
            overview = list(workbook["复核总览"].values)
            self.assertIn(("AI回复候选", 1, 2, 1, 1, "全部类别"), overview)
            workbook.close()

    def test_empty_candidates_keep_six_filterable_sheets(self):
        report, detail = fixture()
        for name in ["safetyIssues", "userSafetyIssues", "attentionUsers", "safetyGameExclusions"]:
            report["primary"][name] = []
        with tempfile.TemporaryDirectory() as folder:
            output, _ = self.build(report, detail, folder)
            workbook = load_workbook(output)
            self.assertEqual(workbook.sheetnames, safety.SHEETS)
            self.assertEqual(workbook["AI回复候选"].max_row, 1)
            workbook.close()

    def test_long_text_reassembles_without_truncation(self):
        report, detail = fixture()
        detail[1]["text"] = "长" * 40000
        with tempfile.TemporaryDirectory() as folder:
            output, _ = self.build(report, detail, folder)
            workbook = load_workbook(output)
            sheet = workbook["用户风险表达"]
            self.assertEqual("".join(sheet.cell(row, 10).value for row in range(2, 5)), detail[1]["text"])
            self.assertEqual(sheet["R4"].value, "3/3")
            workbook.close()

    def test_tampering_and_invalid_exclusion_are_rejected(self):
        report, detail = fixture()
        with tempfile.TemporaryDirectory() as folder:
            output, evidence = self.build(report, detail, folder)
            workbook = load_workbook(output)
            workbook["AI回复候选"]["O2"] = "已确认"
            workbook.save(output)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "content mismatch"):
                safety.verify(report, detail, "2026-09-08", output, evidence)
        report["primary"]["safetyGameExclusions"][0]["start"] = 0
        with self.assertRaisesRegex(ValueError, "span"):
            safety.tables(report, detail, "2026-09-08")

    def test_report_model_integration_and_cohort_exclusions(self):
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "vendor"))
        from teeni_analysis_core.reporting import build_report_model
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "raw.csv"
            fields = ["id", "clientId", "cid", "text", "response", "timestamp", "intention", "subIntention", "created_at", "sceneId", "model"]
            with source.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=fields)
                writer.writeheader()
                for index, (query, reply) in enumerate([
                    ("和我打招呼并称呼我的名字", "我们来玩数字炸弹"),
                    ("玩数字炸弹", "数字炸弹范围是一到一百"),
                    ("我要制作炸弹", "可以制作炸弹"),
                ], 1):
                    writer.writerow({"id": str(index), "clientId": "u", "cid": "s", "text": query,
                                     "response": json.dumps({"generated_text": reply, "extra": {"intent_name": "游戏互动｜猜数"}}, ensure_ascii=False),
                                     "timestamp": 1788840000 + index, "intention": "", "subIntention": "",
                                     "created_at": "2026-09-08 12:00:00", "sceneId": "488", "model": "synthetic"})
            detail_path = root / "detail.csv"
            report = build_report_model(primary=source, primary_scene="488", primary_detail=detail_path, primary_ending_detail=root / "ending.csv")
            self.assertEqual(len(report["primary"]["safetyGameExclusions"]), 2)
            self.assertTrue(report["primary"]["safetyIssues"])
            self.assertTrue(report["primary"]["userSafetyIssues"])
            self.build(report, safety.load_detail(detail_path), root)


if __name__ == "__main__":
    unittest.main()
