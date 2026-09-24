"""Build and reconcile the local-only daily safety review workbook."""
from __future__ import annotations

import argparse
import csv
import json
import math
import sys
import tempfile
import unicodedata
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

SCHEMA = "teeni-safety-review/1.0.0"
SHEETS = ["复核总览", "AI回复候选", "用户风险表达", "需关注用户", "会话上下文", "游戏排除记录"]
MAX_PART = 16000  # Also below Excel's limit for text made entirely of surrogate pairs.


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def load_detail(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def read_input(report_path, detail_path, data_date):
    if date.fromisoformat(data_date).isoformat() != data_date:
        raise ValueError("data date must be YYYY-MM-DD")
    report = load_json(report_path)
    detail = load_detail(detail_path)
    if len(detail) != report["primary"]["summary"]["rows"]:
        raise ValueError("safety review detail count differs from report")
    return report, detail


def key(row):
    return str(row["clientId"]), str(row["cid"]), str(row["id"])


def issue_key(row):
    return str(row["client_id"]), str(row["cid"]), str(row["record_id"])


def _parts(row):
    count = max([math.ceil(len(value) / MAX_PART) for value in row if isinstance(value, str)] + [1])
    for part in range(count):
        yield [
            value[part * MAX_PART:(part + 1) * MAX_PART]
            if isinstance(value, str) and len(value) > MAX_PART else value
            for value in row
        ] + [f"{part + 1}/{count}"]


def tables(report, detail, data_date):
    primary = report["primary"]
    product = {"488": "M1", "904": "M2", "901": "M2"}.get(primary["sceneId"], primary["sceneId"])
    records = {key(row): row for row in detail}
    if len(records) != len(detail):
        raise ValueError("duplicate safety detail identity")
    ai, user = primary["safetyIssues"], primary["userSafetyIssues"]
    exclusions = primary.get("safetyGameExclusions", [])
    for item in [*ai, *user, *exclusions]:
        if issue_key(item) not in records:
            raise ValueError("safety item has no matching source detail")
    selected = {issue_key(item)[:2] for item in [*ai, *user, *exclusions]}
    context_rows = sorted(
        [row for row in detail if key(row)[:2] in selected],
        key=lambda row: (row["clientId"], row["cid"], int(row["turn_index"])),
    )
    context_header = ["账号用户ID", "会话ID", "记录ID", "来源行", "会话轮次", "时间", "用户原文", "AI回复", "AI风险类别", "用户风险类别", "游戏排除命中数", "文本分段"]
    exclusion_counts = Counter(issue_key(item) for item in exclusions)
    context, links = [], {}
    for row in context_rows:
        links[key(row)] = len(context) + 2
        context.extend(_parts([
            row["clientId"], row["cid"], row["id"], int(row["source_row"]), int(row["turn_index"]),
            row["created_at"] or row["timestamp"], row["text"], row["ai_text"],
            row["safety_signals"], row["user_safety_signals"], exclusion_counts[key(row)],
        ]))
    candidate_header = ["账号用户ID", "会话ID", "记录ID", "来源行", "会话轮次", "时间", "风险类别", "严重度", "命中词", "用户原文", "AI回复", "判定依据", "复核优先级", "复核状态", "复核结论", "复核备注", "会话上下文", "文本分段"]

    def candidates(items):
        rows = []
        for item in items:
            source = records[issue_key(item)]
            rows.extend(_parts([
                item["client_id"], item["cid"], item["record_id"], int(source["source_row"]),
                int(source["turn_index"]), item["source_time"], item["category"], item["severity"], item["evidence"],
                source["text"], source["ai_text"], item["reason"], item.get("reviewPriority", ""),
                "待复核", "", "", f"'会话上下文'!A{links[issue_key(item)]}",
            ]))
        return rows

    attention_by_id = {item["clientId"]: item for item in primary["attentionUsers"]}
    attention = []
    ai_by_user, user_by_user = defaultdict(list), defaultdict(list)
    for item in ai:
        ai_by_user[item["client_id"]].append(item)
    for item in user:
        user_by_user[item["client_id"]].append(item)
    for client in sorted(set(ai_by_user) | set(user_by_user), key=lambda value: (attention_by_id.get(value, {}).get("reviewPriority", "P3"), value)):
        ai_items, user_items = ai_by_user[client], user_by_user[client]
        entry = attention_by_id.get(client, {})
        attention.extend(_parts([
            client, entry.get("reviewPriority", ""), len({issue_key(item) for item in ai_items}), len(ai_items),
            len({issue_key(item) for item in user_items}), len(user_items),
            len({item["cid"] for item in ai_items + user_items}),
            "|".join(sorted({item["category"] for item in ai_items})),
            "|".join(sorted({item["category"] for item in user_items})),
            "|".join(sorted({item["cid"] for item in ai_items + user_items})),
            entry.get("priorityReason", "仅AI回复候选，不据此评定用户风险"), "待复核", "", "",
        ]))
    exclusion_rows = []
    for item in exclusions:
        source = records[issue_key(item)]
        side_text = source["text"] if item["side"] == "user" else source["ai_text"]
        start, end = int(item["start"]), int(item["end"])
        if not (0 <= start < end <= len(side_text)) or side_text[start:end] != item["matched_token"]:
            raise ValueError("game exclusion span does not match source text")
        exclusion_rows.extend(_parts([
            item["client_id"], item["cid"], item["record_id"], int(source["source_row"]), int(source["turn_index"]),
            "用户" if item["side"] == "user" else "AI", item["matched_token"], start, end,
            item.get("context_record_id", ""), item["context_evidence"], item["reason"], source["text"], source["ai_text"],
            f"'会话上下文'!A{links[issue_key(item)]}",
        ]))
    overview = [
        ["业务日期", data_date, "产品", product, "", ""],
        ["分析版本", report["coreVersion"], "规则版本", report["rulesVersion"], "", ""],
        ["说明", "自动候选均待复核，不能视为已确认安全问题", "", "", "", ""],
        ["计数口径", "类别命中数允许同一记录多类别；记录数按账号、会话、记录去重", "", "", "", ""],
        ["原文规则", "超长文本按文本分段列连续保存，按序拼接对应长文本列可还原", "", "", "", ""],
    ]
    for label, items in [("AI回复候选", ai), ("用户风险表达", user)]:
        overview.append([label, len({issue_key(item) for item in items}), len(items), len({issue_key(item)[:2] for item in items}), len({item["client_id"] for item in items}), "全部类别"])
        for category in sorted({item["category"] for item in items}):
            group = [item for item in items if item["category"] == category]
            overview.append([label, len({issue_key(item) for item in group}), len(group), len({issue_key(item)[:2] for item in group}), len({item["client_id"] for item in group}), category])
    for side, label in [("ai", "AI游戏排除"), ("user", "用户游戏排除")]:
        group = [item for item in exclusions if item["side"] == side]
        overview.append([label, len({issue_key(item) for item in group}), len(group), len({issue_key(item)[:2] for item in group}), len({item["client_id"] for item in group}), "具体词语命中数"])
    return {
        SHEETS[0]: (["项目", "记录数或说明", "类别或词语命中数", "涉及会话数", "涉及用户数", "风险分类"], overview),
        SHEETS[1]: (candidate_header, candidates(ai)),
        SHEETS[2]: (candidate_header, candidates(user)),
        SHEETS[3]: (["账号用户ID", "用户复核优先级", "AI候选记录数", "AI类别命中数", "用户风险记录数", "用户类别命中数", "涉及会话数", "AI风险类别", "用户风险类别", "涉及会话ID", "优先级依据", "复核状态", "复核结论", "复核备注", "文本分段"], attention),
        SHEETS[4]: (context_header, context),
        SHEETS[5]: (["账号用户ID", "会话ID", "记录ID", "来源行", "会话轮次", "文本侧", "排除命中词", "命中开始位置", "命中结束位置", "语境来源记录ID", "语境依据", "排除原因", "用户原文", "AI回复", "会话上下文", "文本分段"], exclusion_rows),
    }


def build(report, detail, data_date, output, exclusions_path):
    import xlsxwriter
    expected = tables(report, detail, data_date)
    workbook = xlsxwriter.Workbook(output, {"strings_to_formulas": False, "strings_to_urls": False})
    header_format = workbook.add_format({"bold": True, "bg_color": "#185A65", "font_color": "#FFFFFF", "font_name": "Microsoft YaHei", "text_wrap": True, "valign": "vcenter", "border": 1, "border_color": "#D7E0E2"})
    body_format = workbook.add_format({"font_name": "Microsoft YaHei", "font_size": 10, "text_wrap": True, "valign": "top", "border": 1, "border_color": "#E3E8E9"})
    link_format = workbook.add_format({"font_color": "#0563C1", "underline": 1, "valign": "top", "font_name": "Microsoft YaHei"})
    for name, (headers, rows) in expected.items():
        if len(rows) + 1 > 1048576:
            raise ValueError(f"Excel sheet row limit exceeded: {name}; no records were truncated")
        sheet = workbook.add_worksheet(name)
        sheet.freeze_panes(1, 0)
        sheet.autofilter(0, 0, max(1, len(rows)), len(headers) - 1)
        sheet.set_row(0, 40, header_format)
        sheet.write_row(0, 0, headers, header_format)
        widths = []
        for col, title in enumerate(headers):
            width = 54 if title in {"用户原文", "AI回复", "语境依据", "判定依据", "排除原因", "记录数或说明"} else 24 if "ID" in title or "依据" in title else 18
            widths.append(width)
            sheet.set_column(col, col, width, body_format)
        for index, row in enumerate(rows, 1):
            lines = max(sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(char) in "WF" else 1 for char in line) / max(1, widths[col] - 2))) for line in str(value).split("\n")) for col, value in enumerate(row))
            sheet.set_row(index, min(409, max(45, lines * 15 + 10)))
            for col, value in enumerate(row):
                if headers[col] == "会话上下文":
                    result = sheet.write_url(index, col, "internal:" + value, link_format, "查看完整会话")
                else:
                    result = sheet.write(index, col, value, body_format)
                if result != 0:
                    raise ValueError(f"Excel write failure {result}: {name}/{index + 1}/{col + 1}")
        sheet.set_landscape()
        sheet.repeat_rows(0)
    workbook.close()
    Path(exclusions_path).write_text(json.dumps({
        "schemaVersion": SCHEMA, "dataDate": data_date,
        "coreVersion": report["coreVersion"], "rulesVersion": report["rulesVersion"],
        "sceneId": report["primary"]["sceneId"], "exclusions": report["primary"].get("safetyGameExclusions", []),
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return verify(report, detail, data_date, output, exclusions_path)


def verify(report, detail, data_date, output, exclusions_path):
    from openpyxl import load_workbook
    expected = tables(report, detail, data_date)
    exclusions = load_json(exclusions_path)
    if exclusions != {"schemaVersion": SCHEMA, "dataDate": data_date, "coreVersion": report["coreVersion"], "rulesVersion": report["rulesVersion"], "sceneId": report["primary"]["sceneId"], "exclusions": report["primary"].get("safetyGameExclusions", [])}:
        raise ValueError("safety exclusion evidence differs from recomputed report")
    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        if workbook.sheetnames != SHEETS:
            raise ValueError("safety workbook sheet contract mismatch")
        counts = {}
        for name, (headers, rows) in expected.items():
            sheet = workbook[name]
            if sheet.freeze_panes != "A2" or not sheet.auto_filter.ref:
                raise ValueError(f"missing safety freeze/filter: {name}")
            if sheet.max_row != len(rows) + 1 or sheet.max_column != len(headers):
                raise ValueError(f"safety workbook dimensions mismatch: {name}")
            for index, row in enumerate([headers, *rows], 1):
                for col, value in enumerate(row, 1):
                    cell = sheet.cell(index, col)
                    if cell.data_type in {"f", "e"}:
                        raise ValueError(f"unsafe formula/error cell: {name}/{cell.coordinate}")
                    if index > 1 and headers[col - 1] == "会话上下文":
                        if cell.value != "查看完整会话" or not cell.hyperlink or cell.hyperlink.location != value:
                            raise ValueError(f"safety context link mismatch: {name}/{cell.coordinate}")
                    elif (cell.value if cell.value is not None else "") != value:
                        raise ValueError(f"safety workbook content mismatch: {name}/{cell.coordinate}")
            counts[name] = len(rows)
        return {"verified": True, "schemaVersion": SCHEMA, "dataDate": data_date, "sheets": counts, "formulaErrors": 0}
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["build", "verify"])
    for flag in ["detail", "data-date", "output", "exclusions"]:
        parser.add_argument("--" + flag, required=True)
    for flag in ["report", "source", "rules", "scene"]:
        parser.add_argument("--" + flag)
    args = parser.parse_args()
    if args.report:
        report, detail = read_input(args.report, args.detail, args.data_date)
    elif args.command == "verify" and all([args.source, args.rules, args.scene]):
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "vendor"))
        from teeni_analysis_core.reporting import build_report_model
        with tempfile.TemporaryDirectory(prefix="teeni-safety-reconcile-") as temp:
            report = build_report_model(
                primary=args.source, primary_scene=args.scene,
                primary_detail=Path(temp) / "detail.csv",
                primary_ending_detail=Path(temp) / "ending.csv", rules_path=args.rules,
            )
        detail = load_detail(args.detail)
        if len(detail) != report["primary"]["summary"]["rows"]:
            raise ValueError("safety review detail count differs from source")
    else:
        parser.error("build requires --report; verify requires --report or --source/--rules/--scene")
    result = (build if args.command == "build" else verify)(report, detail, args.data_date, args.output, args.exclusions)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
