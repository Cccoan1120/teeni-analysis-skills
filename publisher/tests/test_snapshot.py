import json
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from publisher.privacy import AggregatePrivacyError, assert_aggregate_only
from publisher.snapshot import SnapshotContractError, build_snapshot, sha256_file, validate_snapshot_payload, _validate_risk_summary


ENDING_LABELS = [
    "AI回复失败",
    "纠正、误解或负面摩擦",
    "设备指令或明显话题切换",
    "明确结束或停止",
    "简短确认或自然收尾",
    "实质对话仍在进行但中断",
    "信息不足或无法判断",
]


class SnapshotFixture:
    def __init__(self, root: Path):
        self.root = root
        self.workbook = root / "base.xlsx"
        self.manifest = root / "manifest.json"
        self.source = root / "source.csv"
        self.detail = root / "detail.csv"
        self.ending = root / "ending.csv"

    def create(
        self,
        *,
        intent_header_row: int = 9,
        omit_sheet_dimensions: bool = False,
        invalid_ending_intent_total: bool = False,
        invalid_comparison_delta: bool = False,
        extra_comparison_ending_count: int | None = None,
        scene_id: str = "488",
        source_scene_id: str | None = None,
    ) -> None:
        self.source.write_bytes(b"source-package")
        self.detail.write_bytes(b"aggregate-test-detail")
        self.ending.write_bytes(b"aggregate-test-ending")

        book = Workbook()
        book.remove(book.active)
        for name in (
            "结论总览",
            "剔除单轮开场白后分析",
            "剔除无效轮次后分析",
            "五轮以上结束分析",
            "五轮以上结束分析（不剔除无效）",
            "有效对话子意图对比",
            "真实意图分析",
            "年龄性别概览",
            "所在地分析",
            "星座分析",
            "新版问题总览",
            "新版开场分析",
            "分析口径",
        ):
            book.create_sheet(name)

        overview = book["结论总览"]
        overview["B5"], overview["B6"], overview["B7"] = 1000, 200, 80
        overview["F5"] = 0.75
        overview["J5"], overview["K5"], overview["L5"] = 25, 0.025, "可评分AI回复 1000"
        overview["J6"], overview["K6"], overview["L6"] = 4, 0.004, "可评分AI回复 1000"
        overview["J7"], overview["K7"], overview["L7"] = 3, 0.003, "可审核输入 1000"
        overview["J8"], overview["K8"], overview["L8"] = 2, 0.025, "可审核用户 80"
        overview["J9"], overview["J10"] = "1 / 0 / 1", 0

        engaged = book["剔除单轮开场白后分析"]
        for cell, value in {
            "B6": 200,
            "C6": 150,
            "B8": 5.0,
            "C8": 6.0,
            "B9": 2,
            "C9": 3,
            "B10": 0.3,
            "C10": 0.4,
        }.items():
            engaged[cell] = value

        net = book["剔除无效轮次后分析"]
        for cell, value in {"C6": 120, "C8": 7.0, "C9": 4, "C10": 0.45}.items():
            net[cell] = value

        for sheet_name in ("五轮以上结束分析", "五轮以上结束分析（不剔除无效）"):
            sheet = book[sheet_name]
            sheet["B4"] = 50 if sheet_name == "五轮以上结束分析" else 60
            for offset, label in enumerate(ENDING_LABELS, start=7):
                sheet[f"A{offset}"], sheet[f"B{offset}"], sheet[f"C{offset}"] = label, 5, 1 / 7
                sheet[f"D{offset}"], sheet[f"E{offset}"], sheet[f"F{offset}"] = 8.5, 0.2, 0

        raw = book["五轮以上结束分析（不剔除无效）"]
        for cell, value in {
            "B18": 20,
            "D18": 18,
            "F18": 8,
            "H18": 7,
            "B19": 0.4,
            "D19": 7 / 18,
            "F19": 300,
        }.items():
            raw[cell] = value
        delay_labels = [
            "1分钟内",
            "1-2分钟",
            "2-3分钟",
            "3-5分钟",
            "5-10分钟",
            "10-30分钟",
            "30-60分钟",
            "1-3小时",
            "3-6小时",
            "6-24小时",
            "24小时以上",
            "截至导出未观察到下一会话",
        ]
        for row, label in enumerate(delay_labels, start=28):
            raw[f"A{row}"], raw[f"D{row}"], raw[f"E{row}"] = label, 1, 0.05
            if row < 39:
                raw[f"F{row}"], raw[f"G{row}"] = 0.1, 20

        ending_intent_fixtures = (
            (book["五轮以上结束分析"], 15, (("日常社交", 30), ("情感表达", 20)), (("日常生活、话题闲聊", 35), ("负向情绪", 15))),
            (raw, 43, (("日常社交", 40), ("未返回", 20)), (("日常生活、话题闲聊", 45), ("未返回", 15))),
        )
        for sheet, header_row, primary_rows, secondary_rows in ending_intent_fixtures:
            sheet.cell(row=header_row, column=1, value="真实主意图")
            sheet.cell(row=header_row, column=2, value="数量")
            sheet.cell(row=header_row, column=3, value="占比")
            sheet.cell(row=header_row, column=4, value="真实子意图")
            sheet.cell(row=header_row, column=5, value="数量")
            sheet.cell(row=header_row, column=6, value="占比")
            cohort_sessions = sheet["B4"].value
            for offset, (label, count) in enumerate(primary_rows, start=1):
                if invalid_ending_intent_total and header_row == 15 and offset == 1:
                    count += 1
                sheet.cell(row=header_row + offset, column=1, value=label)
                sheet.cell(row=header_row + offset, column=2, value=count)
                sheet.cell(row=header_row + offset, column=3, value=count / cohort_sessions)
            for offset, (label, count) in enumerate(secondary_rows, start=1):
                sheet.cell(row=header_row + offset, column=4, value=label)
                sheet.cell(row=header_row + offset, column=5, value=count)
                sheet.cell(row=header_row + offset, column=6, value=count / cohort_sessions)

        intent = book["真实意图分析"]
        intent["B5"], intent["B6"] = 900, 100
        intent.cell(row=intent_header_row, column=1, value="主意图")
        intent.cell(row=intent_header_row, column=2, value="数量")
        intent.cell(row=intent_header_row, column=3, value="占比")
        intent.cell(row=intent_header_row, column=4, value="子意图")
        intent.cell(row=intent_header_row, column=5, value="数量")
        intent.cell(row=intent_header_row, column=6, value="占比")
        intent.cell(row=intent_header_row + 1, column=1, value="日常社交")
        intent.cell(row=intent_header_row + 1, column=2, value=500)
        intent.cell(row=intent_header_row + 1, column=3, value=0.5)
        intent.cell(row=intent_header_row + 1, column=4, value="日常生活、话题闲聊")
        intent.cell(row=intent_header_row + 1, column=5, value=450)
        intent.cell(row=intent_header_row + 1, column=6, value=0.45)

        comparison = book["有效对话子意图对比"]
        comparison.append(["有效对话子意图对比"])
        comparison.append(["比较全部净有效轮次与五轮以上有效末轮；百分点差=末轮占比-全部占比。"])
        comparison.append([])
        comparison.append(["真实子意图", "全部净有效轮次", "全部占比", "五轮末轮会话", "末轮占比", "百分点差"])
        comparison_rows = [
            ("日常生活、话题闲聊", 70, 35),
            ("负向情绪", 30, 15),
        ]
        if extra_comparison_ending_count is not None:
            comparison_rows.append(("知识问答", 0, extra_comparison_ending_count))
        all_total = sum(row[1] for row in comparison_rows)
        ending_total = sum(row[2] for row in comparison_rows)
        for index, (label, all_count, ending_count) in enumerate(comparison_rows):
            all_share = all_count / all_total
            ending_share = ending_count / ending_total
            share_delta = ending_share - all_share
            if invalid_comparison_delta and index == 0:
                share_delta += 0.01
            comparison.append([label, all_count, all_share, ending_count, ending_share, share_delta])

        profile = book["年龄性别概览"]
        profile["B4"], profile["D4"], profile["F4"] = 80, 75, 0.9375
        profile["A14"], profile["B14"] = "年龄段", "性别"
        values = ["3-4岁", "女", 40, 90, 600, 6.6667, 30, 1 / 3, "可描述"]
        for column, value in enumerate(values, start=1):
            profile.cell(row=15, column=column, value=value)

        location = book["所在地分析"]
        location["B4"], location["D4"], location["F4"] = 80, 50, 0.625
        location["A6"], location["B6"] = "解析状态", "用户数"
        location["A7"], location["B7"] = "正常", 50
        location["A8"], location["B8"] = "缺失", 30
        location["A10"] = "城市"
        for row, values in enumerate(
            [
                ("北京", 30, 0.6, 50, 300, 6, 20, 0.4, "可描述"),
                ("阿克苏地区", 20, 0.4, 30, 150, 5, 10, 1 / 3, "样本不足"),
            ],
            start=11,
        ):
            for column, value in enumerate(values, start=1):
                location.cell(row=row, column=column, value=value)

        constellation = book["星座分析"]
        constellation["B4"], constellation["D4"], constellation["F4"] = 80, 48, 0.6
        constellation["A6"], constellation["B6"] = "解析状态", "用户数"
        constellation["A7"], constellation["B7"] = "正常", 48
        constellation["A8"], constellation["B8"] = "缺失", 32
        constellation["A10"] = "星座"
        labels = ["摩羯座", "水瓶座", "双鱼座", "白羊座", "金牛座", "双子座", "巨蟹座", "狮子座", "处女座", "天秤座", "天蝎座", "射手座"]
        for row, label in enumerate(labels, start=11):
            values = (label, 4, 1 / 12, 8, 48, 6, 2, 0.25, "样本不足")
            for column, value in enumerate(values, start=1):
                constellation.cell(row=row, column=column, value=value)

        problems = book["新版问题总览"]
        for row in range(5, 28):
            problems[f"B{row}"], problems[f"C{row}"] = f"聚合类别{row}", "中"
            problems[f"D{row}"], problems[f"E{row}"], problems[f"F{row}"] = 1, 1, 1
            problems[f"G{row}"], problems[f"H{row}"] = 0.001, "聚合分母 1000"

        opening = book["新版开场分析"]
        opening["A4"], opening["B4"], opening["C4"], opening["D4"] = "难度", "曝光", "开口", "开口率"
        for row, values in enumerate(
            [("中", 100, 75), ("低", 50, 25), ("高", 25, 10)],
            start=5,
        ):
            opening[f"A{row}"], opening[f"B{row}"], opening[f"C{row}"] = values
            opening[f"D{row}"] = values[2] / values[1]

        method = book["分析口径"]
        method["B6"], method["B7"], method["B8"] = "2.5.1", "13.0.0", "teeni-base-detail/1.2.0"
        book.save(self.workbook)
        if omit_sheet_dimensions:
            rewritten = self.workbook.with_suffix(".rewritten.xlsx")
            namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            with zipfile.ZipFile(self.workbook, "r") as source, zipfile.ZipFile(rewritten, "w") as target:
                for info in source.infolist():
                    data = source.read(info.filename)
                    if info.filename.startswith("xl/worksheets/sheet") and info.filename.endswith(".xml"):
                        root = ET.fromstring(data)
                        dimension = root.find(f"{namespace}dimension")
                        if dimension is not None:
                            root.remove(dimension)
                            data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                    target.writestr(info, data)
            rewritten.replace(self.workbook)

        payload = {
            "schemaVersion": "teeni-base-bundle-manifest/1.0.0",
            "contractVersion": "teeni-base-detail/1.2.0",
            "coreVersion": "2.5.1",
            "rulesVersion": "13.0.0",
            "mode": "primary",
            "sceneId": scene_id,
            "sources": [
                {
                    "role": "primary",
                    "file": self.source.name,
                    "sha256": sha256_file(self.source),
                    "rows": 1000,
                    "sceneId": source_scene_id or scene_id,
                }
            ],
            "outputs": [
                {"role": "workbook", "file": self.workbook.name, "sha256": sha256_file(self.workbook)},
                {"role": "primary_detail", "file": self.detail.name, "sha256": sha256_file(self.detail), "rows": 1000},
                {"role": "ending_detail", "file": self.ending.name, "sha256": sha256_file(self.ending), "rows": 50},
            ],
        }
        self.manifest.write_text(json.dumps(payload), encoding="utf-8")


class SnapshotTests(unittest.TestCase):
    def test_builds_aggregate_snapshot_from_verified_package(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create()
            snapshot = build_snapshot(fixture.workbook, fixture.manifest, date(2026, 8, 17))
            self.assertEqual(snapshot["dataDate"], "2026-08-17")
            self.assertEqual(snapshot["schemaVersion"], "teeni-dashboard-snapshot/1.2.0")
            self.assertEqual(snapshot["productVersion"], "M1")
            self.assertEqual(snapshot["sceneId"], "488")
            self.assertEqual(snapshot["metrics"]["participation"]["sessions"], 150)
            self.assertEqual(snapshot["metrics"]["opening"]["exposures"], 175)
            self.assertEqual(snapshot["metrics"]["opening"]["opened"], 110)
            self.assertAlmostEqual(snapshot["metrics"]["opening"]["rate"], 110 / 175)
            self.assertEqual(snapshot["metrics"]["reopen"]["fiveMinuteSessions"], 8)
            self.assertEqual(snapshot["metrics"]["endings"]["validIntents"]["primary"][0]["label"], "日常社交")
            self.assertEqual(snapshot["metrics"]["endings"]["validIntents"]["primary"][0]["count"], 30)
            self.assertEqual(snapshot["metrics"]["endings"]["validIntents"]["secondary"][0]["label"], "日常生活、话题闲聊")
            self.assertEqual(snapshot["metrics"]["endings"]["rawIntents"]["primary"][1]["label"], "未返回")
            self.assertEqual(snapshot["metrics"]["endings"]["rawIntents"]["secondary"][0]["count"], 45)
            comparison = snapshot["metrics"]["validSecondaryIntentComparison"]
            self.assertEqual(comparison["allNetValidTurns"], 100)
            self.assertEqual(comparison["endingSessions"], 50)
            self.assertEqual(comparison["rows"][0]["endingCount"], 35)
            self.assertEqual(snapshot["metrics"]["riskSummary"]["priority"]["p0"], 1)
            self.assertEqual(snapshot["metrics"]["locations"]["items"][0]["label"], "北京")
            self.assertEqual(snapshot["metrics"]["locations"]["items"][1]["label"], "阿克苏地区")
            self.assertEqual(len(snapshot["metrics"]["constellations"]["items"]), 12)
            self.assertAlmostEqual(snapshot["metrics"]["coverage"]["locationRate"], 0.625)
            self.assertAlmostEqual(snapshot["metrics"]["coverage"]["constellationRate"], 0.6)
            serialized = json.dumps(snapshot, ensure_ascii=False)
            self.assertNotIn("clientId", serialized)
            self.assertNotIn("ai_text", serialized)
            self.assertNotIn("response", serialized)

    def test_rejects_ending_intent_totals_that_do_not_match_cohort(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create(invalid_ending_intent_total=True)
            with self.assertRaisesRegex(SnapshotContractError, "valid ending primary intent count mismatch"):
                build_snapshot(fixture.workbook, fixture.manifest, date(2026, 8, 30))

    def test_rejects_invalid_valid_secondary_intent_comparison_delta(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create(invalid_comparison_delta=True)
            with self.assertRaisesRegex(SnapshotContractError, "valid secondary intent delta row 5 mismatch"):
                build_snapshot(fixture.workbook, fixture.manifest, date(2026, 8, 31))

    def test_accepts_zero_count_secondary_intent_comparison_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create(extra_comparison_ending_count=0)
            snapshot = build_snapshot(fixture.workbook, fixture.manifest, date(2026, 9, 1))
            self.assertEqual(snapshot["metrics"]["validSecondaryIntentComparison"]["rows"][-1]["endingCount"], 0)

    def test_rejects_nonzero_secondary_intent_comparison_rows_missing_from_endings(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create(extra_comparison_ending_count=1)
            with self.assertRaisesRegex(SnapshotContractError, "valid secondary intent ending rows mismatch"):
                build_snapshot(fixture.workbook, fixture.manifest, date(2026, 9, 1))

    def test_current_snapshot_without_ending_intents_remains_valid(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create()
            snapshot = build_snapshot(fixture.workbook, fixture.manifest, date(2026, 8, 30))
            snapshot["metrics"]["endings"].pop("validIntents")
            snapshot["metrics"]["endings"].pop("rawIntents")
            validate_snapshot_payload(snapshot)

    def test_accepts_current_intent_header_row(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create(intent_header_row=8)
            snapshot = build_snapshot(fixture.workbook, fixture.manifest, date(2026, 8, 24))
            self.assertEqual(snapshot["metrics"]["intents"]["primary"][0]["label"], "日常社交")
            self.assertEqual(snapshot["metrics"]["intents"]["primary"][0]["count"], 500)
            self.assertEqual(snapshot["metrics"]["intents"]["secondary"][0]["label"], "日常生活、话题闲聊")

    def test_accepts_streamed_sheets_without_dimension_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create(omit_sheet_dimensions=True)
            snapshot = build_snapshot(fixture.workbook, fixture.manifest, date(2026, 8, 25))
            self.assertEqual([row["label"] for row in snapshot["metrics"]["locations"]["items"]], ["北京", "阿克苏地区"])
            self.assertEqual(len(snapshot["metrics"]["constellations"]["items"]), 12)

    def test_rejects_wrong_scene(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create()
            manifest = json.loads(fixture.manifest.read_text(encoding="utf-8"))
            manifest["sceneId"] = "901"
            fixture.manifest.write_text(json.dumps(manifest), encoding="utf-8")
            with self.assertRaisesRegex(SnapshotContractError, "scene 488 or 904"):
                build_snapshot(fixture.workbook, fixture.manifest, date(2026, 8, 17))

    def test_builds_m2_snapshot_from_scene_904_package(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create(scene_id="904")
            snapshot = build_snapshot(fixture.workbook, fixture.manifest, date(2026, 9, 1))
            self.assertEqual(snapshot["productVersion"], "M2")
            self.assertEqual(snapshot["sceneId"], "904")

    def test_rejects_m2_snapshot_before_history_start(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create(scene_id="904")
            with self.assertRaisesRegex(SnapshotContractError, "start on 2026-09-01"):
                build_snapshot(fixture.workbook, fixture.manifest, date(2026, 8, 31))

    def test_rejects_manifest_source_scene_mismatch(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create(scene_id="904", source_scene_id="488")
            with self.assertRaisesRegex(SnapshotContractError, "source scene"):
                build_snapshot(fixture.workbook, fixture.manifest, date(2026, 9, 1))

    def test_rejects_snapshot_product_scene_mismatch(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create(scene_id="904")
            snapshot = build_snapshot(fixture.workbook, fixture.manifest, date(2026, 9, 1))
            snapshot["productVersion"] = "M1"
            with self.assertRaisesRegex(SnapshotContractError, "does not match scene"):
                validate_snapshot_payload(snapshot)

    def test_rejects_workbook_hash_mismatch(self):
        with tempfile.TemporaryDirectory() as directory:
            fixture = SnapshotFixture(Path(directory))
            fixture.create()
            with fixture.workbook.open("ab") as stream:
                stream.write(b"changed")
            with self.assertRaisesRegex(SnapshotContractError, "workbook hash mismatch"):
                build_snapshot(fixture.workbook, fixture.manifest, date(2026, 8, 17))

    def test_privacy_validator_rejects_sensitive_shapes(self):
        cases = [
            {"clientId": "abc"},
            {"query": "raw content"},
            {"safe": {"ai_text": "content"}},
            {"safe": {"birthday": "2021-05-28"}},
            {"safe": {"city": "北京"}},
            {"safe": {"constellation": "双子座"}},
            {"safe": "1234567890abcdef1234"},
            {"safe": "manager@example.com"},
            {"safe": "a" * 241},
            {"metrics": {"supplemental": {"quality": {"responseMissing": "raw reply"}}}},
            {"metrics": {"supplemental": {"quality": {"responseMissing": {"numerator": "raw reply", "denominator": 1, "rate": 1}}}}},
            {"metrics": {"supplemental": {"inputQuality": {"missingKeys": {"clientId": "raw identifier"}}}}},
        ]
        for payload in cases:
            with self.subTest(payload=payload):
                with self.assertRaises(AggregatePrivacyError):
                    assert_aggregate_only(payload)

        assert_aggregate_only({"locations": [{"label": "北京", "users": 10}]})
        assert_aggregate_only({"queryCount": 10})
        assert_aggregate_only({"metrics": {"supplemental": {"quality": {"responseMissing": {"numerator": 1, "denominator": 2, "rate": 0.5}}}}})
        assert_aggregate_only(
            {
                "constellations": [{"label": "双子座", "users": 10}],
                "constellationUsers": 10,
                "constellationEligibleUsers": 20,
                "constellationRate": 0.5,
            }
        )

    def test_privacy_validator_accepts_sha256_with_phone_like_digits(self):
        assert_aggregate_only(
            {"workbookSha256": "62b2ade931f91eab55e18240425990e9b4eed1c2312e35e795b142d33f986227"}
        )


class RiskRateContractTests(unittest.TestCase):
    def summary(self, denominator):
        return dict(scoreableAiReplies=denominator, reviewableUserInputs=denominator, reviewableUsers=denominator,
                    qualityCandidates=0, aiSafetyCandidates=0, userRiskCandidates=0, candidateUsers=0,
                    qualityRate=0 if denominator else None, aiSafetyRate=0 if denominator else None,
                    userRiskRate=0 if denominator else None, candidateUserRate=0 if denominator else None)

    def test_no_sample_requires_null_but_observed_zero_is_valid(self):
        _validate_risk_summary(self.summary(0))
        _validate_risk_summary(self.summary(10))
        for key in ('qualityRate', 'aiSafetyRate', 'userRiskRate', 'candidateUserRate'):
            value = self.summary(0)
            value[key] = 0
            with self.subTest(key=key), self.assertRaises(SnapshotContractError):
                _validate_risk_summary(value)

    def test_event_density_can_exceed_one_but_unique_rate_cannot(self):
        value = self.summary(1)
        value.update(aiSafetyCandidates=3, aiSafetyRate=3)
        _validate_risk_summary(value)
        for count, rate in (('qualityCandidates', 'qualityRate'), ('userRiskCandidates', 'userRiskRate'), ('candidateUsers', 'candidateUserRate')):
            value = self.summary(1)
            value.update({count: 2, rate: 2})
            with self.subTest(rate=rate), self.assertRaises(SnapshotContractError):
                _validate_risk_summary(value)


if __name__ == "__main__":
    unittest.main()
