import hashlib
import csv
import json
import re
import math
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from dashboard.product_versions import PRODUCT_VERSION_SCENES, product_version_for_scene

from .privacy import assert_aggregate_only
from .date_contract import CsvDateContractError, validate_csv_data_date
from .session_structure import from_detail, validate_structure


SNAPSHOT_SCHEMA = "teeni-dashboard-snapshot/1.3.0"
PREVIOUS_PRODUCT_SNAPSHOT_SCHEMA = "teeni-dashboard-snapshot/1.2.0"
PREVIOUS_SNAPSHOT_SCHEMA = "teeni-dashboard-snapshot/1.1.0"
LEGACY_SNAPSHOT_SCHEMA = "teeni-dashboard-snapshot/1.0.0"
MANIFEST_SCHEMA = "teeni-base-bundle-manifest/1.0.0"
EXPECTED_CONTRACT = "teeni-base-detail/1.2.0"
EXPECTED_CORE = "2.5.1"
EXPECTED_RULES = "13.0.0"
PREVIOUS_CONTRACT = "teeni-base-detail/1.2.0"
PREVIOUS_CORE = "2.5.0"
PREVIOUS_RULES = "11.0.0"
LEGACY_CONTRACT = "teeni-base-detail/1.1.0"
LEGACY_CORE = "2.4.0"
LEGACY_RULES = "10.0.0"
M2_FIRST_DATA_DATE = date(2026, 9, 1)


class SnapshotContractError(ValueError):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise SnapshotContractError(message)


def _as_int(value: Any, label: str) -> int:
    _require(isinstance(value, (int, float)) and not isinstance(value, bool), f"{label} must be numeric")
    return int(value)


def _as_float(value: Any, label: str) -> float:
    _require(isinstance(value, (int, float)) and not isinstance(value, bool), f"{label} must be numeric")
    return float(value)


def _nullable_float(value: Any, label: str) -> float | None:
    return None if value is None else _as_float(value, label)


def _as_label(value: Any, label: str) -> str:
    _require(isinstance(value, str) and 0 < len(value) <= 80, f"{label} must be a short label")
    _require(bool(value.strip()) and not re.search(r"<[!/]?[A-Za-z][^>]*>", value), f"{label} must be plain text")
    return value.strip()


def _number_from_label(value: Any, label: str) -> int:
    _require(isinstance(value, str), f"{label} must contain an aggregate count")
    matches = re.findall(r"\d+", value)
    _require(bool(matches), f"{label} does not contain an aggregate count")
    return int(matches[-1])


def _validate_manifest(
    manifest: dict[str, Any],
    workbook: Path,
    full_hash_check: bool,
    primary_source: Path | None = None,
) -> tuple[str, str, str, str]:
    _require(manifest.get("schemaVersion") == MANIFEST_SCHEMA, "unsupported manifest schema")
    _require(manifest.get("contractVersion") == EXPECTED_CONTRACT, "unsupported detail contract")
    _require((manifest.get("coreVersion"), manifest.get("rulesVersion")) in {(EXPECTED_CORE, EXPECTED_RULES), ("2.6.0", "14.0.0"), ("2.6.1", "15.0.0"), ("2.7.0", "16.0.0")}, "unsupported core/rules version")
    scene_id = str(manifest.get("sceneId"))
    product_version = product_version_for_scene(scene_id)
    _require(product_version is not None, "dashboard accepts scene 488 or 904 only")
    _require(manifest.get("mode") == "primary", "dashboard accepts primary packages only")

    sources = manifest.get("sources")
    outputs = manifest.get("outputs")
    _require(isinstance(sources, list) and len(sources) == 1, "manifest must contain one primary source")
    _require(isinstance(outputs, list), "manifest outputs are missing")
    if manifest.get("rulesVersion") in {"15.0.0", "16.0.0"}:
        for role in ("safety_workbook", "safety_exclusions"):
            _require(sum(item.get("role") == role for item in outputs) == 1, f"new safety contract requires {role}")
    source = sources[0]
    _require(source.get("role") == "primary", "manifest must identify its primary source")
    _require(str(source.get("sceneId")) == scene_id, "manifest source scene does not match package scene")
    workbook_entries = [item for item in outputs if item.get("role") == "workbook"]
    _require(len(workbook_entries) == 1, "manifest must contain one workbook output")
    workbook_entry = workbook_entries[0]
    _require(workbook.name == workbook_entry.get("file"), "workbook filename does not match manifest")
    _require(sha256_file(workbook) == workbook_entry.get("sha256"), "workbook hash mismatch")

    if primary_source is not None:
        _require(primary_source.name == source.get("file"), "primary source filename does not match manifest")
        _require(primary_source.is_file(), "primary source does not exist")
        _require(sha256_file(primary_source) == source.get("sha256"), "primary source hash mismatch")

    if full_hash_check:
        base_dir = workbook.parent
        for item in [*sources, *outputs]:
            file_name = item.get("file")
            _require(isinstance(file_name, str) and Path(file_name).name == file_name, "manifest filename is unsafe")
            file_path = primary_source if item is source and primary_source is not None else base_dir / file_name
            _require(file_path.is_file(), f"package companion is missing: {file_name}")
            _require(sha256_file(file_path) == item.get("sha256"), f"package hash mismatch: {file_name}")

    return str(source["sha256"]), str(workbook_entry["sha256"]), scene_id, product_version


def _problem_rows(sheet, start: int, end: int) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in range(start, end + 1):
        category = sheet[f"B{row}"].value
        if not category:
            continue
        rows.append(
            {
                "category": _as_label(category, f"problem category row {row}"),
                "severity": _as_label(sheet[f"C{row}"].value, f"problem severity row {row}"),
                "count": _as_int(sheet[f"D{row}"].value, f"problem count row {row}"),
                "sessions": _as_int(sheet[f"E{row}"].value, f"problem sessions row {row}"),
                "users": _as_int(sheet[f"F{row}"].value, f"problem users row {row}"),
                "rate": _as_float(sheet[f"G{row}"].value, f"problem rate row {row}"),
                "denominator": _number_from_label(sheet[f"H{row}"].value, f"problem denominator row {row}"),
            }
        )
    return rows


def _ending_rows(sheet) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in range(7, 14):
        rows.append(
            {
                "category": _as_label(sheet[f"A{row}"].value, f"ending category row {row}"),
                "sessions": _as_int(sheet[f"B{row}"].value, f"ending sessions row {row}"),
                "share": _as_float(sheet[f"C{row}"].value, f"ending share row {row}"),
                "averageDepth": _as_float(sheet[f"D{row}"].value, f"ending depth row {row}"),
                "followupShare": _as_float(sheet[f"E{row}"].value, f"ending followup row {row}"),
                "boundaryCount": _as_int(sheet[f"F{row}"].value, f"ending boundary row {row}"),
            }
        )
    return rows


def _intent_rows(sheet, label_col: str, count_col: str, share_col: str, start: int, end: int) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in range(start, end + 1):
        label = sheet[f"{label_col}{row}"].value
        if not label:
            continue
        rows.append(
            {
                "label": _as_label(label, f"intent label row {row}"),
                "count": _as_int(sheet[f"{count_col}{row}"].value, f"intent count row {row}"),
                "share": _as_float(sheet[f"{share_col}{row}"].value, f"intent share row {row}"),
            }
        )
    return rows


def _ending_intents(sheet, cohort_sessions: int, cohort_label: str) -> dict[str, list[dict[str, Any]]]:
    header = next(
        (
            row
            for row in range(1, 101)
            if [sheet[f"{column}{row}"].value for column in "ABCDEF"]
            == ["真实主意图", "数量", "占比", "真实子意图", "数量", "占比"]
        ),
        None,
    )
    _require(header is not None, f"{cohort_label} ending intent headers are invalid")

    result: dict[str, list[dict[str, Any]]] = {"primary": [], "secondary": []}
    columns = (("primary", "A", "B", "C"), ("secondary", "D", "E", "F"))
    for row in range(header + 1, header + 101):
        labels = [sheet[f"A{row}"].value, sheet[f"D{row}"].value]
        if all(label in (None, "") for label in labels):
            break
        for kind, label_col, count_col, share_col in columns:
            label = sheet[f"{label_col}{row}"].value
            if label in (None, ""):
                continue
            count = _as_int(sheet[f"{count_col}{row}"].value, f"{cohort_label} ending {kind} count row {row}")
            share = _as_float(sheet[f"{share_col}{row}"].value, f"{cohort_label} ending {kind} share row {row}")
            _require(count >= 0, f"{cohort_label} ending {kind} count row {row} must be non-negative")
            _require(abs(share - count / cohort_sessions) < 1e-12, f"{cohort_label} ending {kind} share row {row} mismatch")
            result[kind].append({
                "label": _as_label(label, f"{cohort_label} ending {kind} label row {row}"),
                "count": count,
                "share": share,
            })

    for kind in ("primary", "secondary"):
        _require(bool(result[kind]), f"{cohort_label} ending {kind} intents are missing")
        _require(
            sum(item["count"] for item in result[kind]) == cohort_sessions,
            f"{cohort_label} ending {kind} intent count mismatch",
        )
    return result


def _valid_secondary_intent_comparison(sheet) -> dict[str, Any]:
    _require(
        [sheet[f"{column}4"].value for column in "ABCDEF"]
        == ["真实子意图", "全部净有效轮次", "全部占比", "五轮末轮会话", "末轮占比", "百分点差"],
        "valid secondary intent comparison headers are invalid",
    )
    rows: list[dict[str, Any]] = []
    labels: set[str] = set()
    for row in range(5, 105):
        label = sheet[f"A{row}"].value
        if label in (None, ""):
            break
        intent_label = _as_label(label, f"valid secondary intent comparison label row {row}")
        _require(intent_label not in labels, "valid secondary intent comparison labels must be unique")
        labels.add(intent_label)
        all_count = _as_int(sheet[f"B{row}"].value, f"valid secondary intent all count row {row}")
        all_share = _as_float(sheet[f"C{row}"].value, f"valid secondary intent all share row {row}")
        ending_count = _as_int(sheet[f"D{row}"].value, f"valid secondary intent ending count row {row}")
        ending_share = _as_float(sheet[f"E{row}"].value, f"valid secondary intent ending share row {row}")
        share_delta = _as_float(sheet[f"F{row}"].value, f"valid secondary intent delta row {row}")
        _require(min(all_count, ending_count) >= 0, f"valid secondary intent counts row {row} must be non-negative")
        _require(abs(share_delta - (ending_share - all_share)) < 1e-12, f"valid secondary intent delta row {row} mismatch")
        rows.append({
            "label": intent_label,
            "allCount": all_count,
            "allShare": all_share,
            "endingCount": ending_count,
            "endingShare": ending_share,
            "shareDelta": share_delta,
        })

    _require(bool(rows), "valid secondary intent comparison rows are missing")
    all_total = sum(item["allCount"] for item in rows)
    ending_total = sum(item["endingCount"] for item in rows)
    _require(all_total > 0 and ending_total > 0, "valid secondary intent comparison totals must be positive")
    for index, item in enumerate(rows):
        _require(abs(item["allShare"] - item["allCount"] / all_total) < 1e-12, f"valid secondary intent all share row {index} mismatch")
        _require(abs(item["endingShare"] - item["endingCount"] / ending_total) < 1e-12, f"valid secondary intent ending share row {index} mismatch")
    return {"allNetValidTurns": all_total, "endingSessions": ending_total, "rows": rows}


def _profile_segment(sheet, group_label: str) -> dict[str, Any]:
    eligible_users = _as_int(sheet["B4"].value, f"{group_label} eligible users")
    normal_users = _as_int(sheet["D4"].value, f"{group_label} normal users")
    coverage_rate = _as_float(sheet["F4"].value, f"{group_label} coverage")
    statuses: list[dict[str, Any]] = []
    for row in range(7, 20):
        status = sheet[f"A{row}"].value
        if status in (None, ""):
            break
        statuses.append({
            "status": _as_label(status, f"{group_label} status row {row}"),
            "users": _as_int(sheet[f"B{row}"].value, f"{group_label} status users row {row}"),
        })
    header = next(
        (row for row in range(8, 30) if sheet[f"A{row}"].value == group_label),
        None,
    )
    _require(header is not None, f"{group_label} aggregate header is missing")
    items: list[dict[str, Any]] = []
    for row, values in enumerate(
        sheet.iter_rows(min_row=header + 1, max_col=9, values_only=True),
        start=header + 1,
    ):
        label = values[0]
        if label in (None, ""):
            break
        items.append({
            "label": _as_label(label, f"{group_label} label row {row}"),
            "users": _as_int(values[1], f"{group_label} users row {row}"),
            "userShare": _as_float(values[2], f"{group_label} user share row {row}"),
            "sessions": _as_int(values[3], f"{group_label} sessions row {row}"),
            "netValidTurns": _as_int(values[4], f"{group_label} turns row {row}"),
            "averageDepth": _as_float(values[5], f"{group_label} depth row {row}"),
            "fivePlusSessions": _as_int(values[6], f"{group_label} five-plus row {row}"),
            "fivePlusShare": _as_float(values[7], f"{group_label} five-plus share row {row}"),
            "sampleStatus": _as_label(values[8], f"{group_label} sample status row {row}"),
        })
    return {
        "eligibleUsers": eligible_users,
        "normalUsers": normal_users,
        "coverageRate": coverage_rate,
        "statuses": statuses,
        "items": items,
    }


def _opening_metrics(sheet, current=False) -> dict[str, Any]:
    _require(
        [sheet[f"{column}4"].value for column in "ABCD"] == (["难度", "原始曝光", "后续有记录", "记录接续率"] if current else ["难度", "曝光", "开口", "开口率"]),
        "opening aggregate headers are invalid",
    )
    expected_difficulties = {"中", "低", "高"}
    difficulties: set[str] = set()
    exposures = 0
    opened = 0
    for row in range(5, 8):
        difficulty = _as_label(sheet[f"A{row}"].value, f"opening difficulty row {row}")
        row_exposures = _as_int(sheet[f"B{row}"].value, f"opening exposures row {row}")
        row_opened = _as_int(sheet[f"C{row}"].value, f"opening opened row {row}")
        _require(row_exposures >= 0, f"opening exposures row {row} must be non-negative")
        _require(0 <= row_opened <= row_exposures, f"opening opened row {row} is invalid")
        difficulties.add(difficulty)
        exposures += row_exposures
        opened += row_opened
    _require(difficulties == expected_difficulties, "opening difficulty rows are incomplete")
    return {"exposures": exposures, "opened": opened, "rate": opened / exposures if exposures else None}


def _extract_workbook(workbook: Path, versions=(EXPECTED_CORE, EXPECTED_RULES)) -> dict[str, Any]:
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        return _extract_open_workbook(book, versions)
    finally:
        book.close()


def _extract_open_workbook(book, versions=(EXPECTED_CORE, EXPECTED_RULES)) -> dict[str, Any]:
    required = {
        "结论总览",
        "剔除单轮开场白后分析",
        "剔除无效轮次后分析",
        "五轮以上结束分析",
        "五轮以上结束分析（不剔除无效）",
        "真实意图分析",
        "年龄性别概览",
        "所在地分析",
        "星座分析",
        "新版问题总览",
        "新版开场分析",
        "分析口径",
    }
    _require(required.issubset(set(book.sheetnames)), "workbook is missing required aggregate sheets")

    overview = book["结论总览"]
    engaged = book["剔除单轮开场白后分析"]
    net = book["剔除无效轮次后分析"]
    valid_ending = book["五轮以上结束分析"]
    raw_ending = book["五轮以上结束分析（不剔除无效）"]
    intent = book["真实意图分析"]
    valid_intent_comparison = (
        _valid_secondary_intent_comparison(book["有效对话子意图对比"])
        if "有效对话子意图对比" in book.sheetnames
        else None
    )
    profile = book["年龄性别概览"]
    locations = _profile_segment(book["所在地分析"], "城市")
    constellations = _profile_segment(book["星座分析"], "星座")
    problems = book["新版问题总览"]
    opening = book["新版开场分析"]
    methodology = book["分析口径"]

    _require(methodology["B6"].value == versions[0], "workbook core version mismatch")
    _require(methodology["B7"].value == versions[1], "workbook rules version mismatch")
    _require(methodology["B8"].value == EXPECTED_CONTRACT, "workbook detail contract mismatch")

    priority_text = _as_label(overview["J9"].value, "review priority counts")
    priority_counts = [int(value) for value in re.findall(r"\d+", priority_text)]
    _require(len(priority_counts) == 3, "P0/P1/P2 counts are invalid")
    intent_header = next(
        (
            row
            for row in range(8, 11)
            if [intent[f"{column}{row}"].value for column in "ABCDEF"]
            == ["主意图", "数量", "占比", "子意图", "数量", "占比"]
        ),
        None,
    )
    _require(intent_header is not None, "intent aggregate headers are invalid")

    demographic_rows: list[dict[str, Any]] = []
    demographic_header = next(
        (
            row
            for row in range(10, 25)
            if profile[f"A{row}"].value == "年龄段" and profile[f"B{row}"].value == "性别"
        ),
        None,
    )
    _require(demographic_header is not None, "demographic aggregate header is missing")
    empty_rows = 0
    for row in range(demographic_header + 1, 80):
        if not profile[f"A{row}"].value:
            empty_rows += 1
            if empty_rows >= 3:
                break
            continue
        empty_rows = 0
        demographic_rows.append(
            {
                "ageBand": _as_label(profile[f"A{row}"].value, f"age band row {row}"),
                "gender": _as_label(profile[f"B{row}"].value, f"gender row {row}"),
                "users": _as_int(profile[f"C{row}"].value, f"profile users row {row}"),
                "sessions": _as_int(profile[f"D{row}"].value, f"profile sessions row {row}"),
                "netValidTurns": _as_int(profile[f"E{row}"].value, f"profile turns row {row}"),
                "averageDepth": _as_float(profile[f"F{row}"].value, f"profile depth row {row}"),
                "fivePlusSessions": _as_int(profile[f"G{row}"].value, f"profile five plus row {row}"),
                "fivePlusShare": _as_float(profile[f"H{row}"].value, f"profile share row {row}"),
                "sampleStatus": _as_label(profile[f"I{row}"].value, f"sample status row {row}"),
            }
        )

    delay_buckets: list[dict[str, Any]] = []
    for row in range(28, 40):
        bucket = raw_ending[f"A{row}"].value
        if not bucket:
            continue
        delay_buckets.append(
            {
                "bucket": _as_label(bucket, f"delay bucket row {row}"),
                "sessions": _as_int(raw_ending[f"D{row}"].value, f"delay sessions row {row}"),
                "shareOfTriggers": _as_float(raw_ending[f"E{row}"].value, f"delay trigger share row {row}"),
                "shareOfObserved": None
                if raw_ending[f"F{row}"].value in (None, "")
                else _as_float(raw_ending[f"F{row}"].value, f"delay observed share row {row}"),
                "observableSample": None
                if raw_ending[f"G{row}"].value in (None, "")
                else _as_int(raw_ending[f"G{row}"].value, f"delay sample row {row}"),
            }
        )

    parsed_intents = _as_int(intent["B5"].value, "parsed intents")
    missing_intents = _as_int(intent["B6"].value, "missing intents")

    valid_ending_sessions = _as_int(valid_ending["B4"].value, "valid ending cohort")
    raw_ending_sessions = _as_int(raw_ending["B4"].value, "raw ending cohort")

    result = {
        "volume": {
            "turns": _as_int(overview["B5"].value, "total turns"),
            "sessions": _as_int(overview["B6"].value, "total sessions"),
            "users": _as_int(overview["B7"].value, "total users"),
        },
        "participation": {
            "sessions": _as_int(engaged["C6"].value, "participating sessions"),
            "eligibleSessions": _as_int(engaged["B6"].value, "eligible sessions"),
            "rate": _as_float(overview["F5"].value, "participation rate"),
        },
        "opening": _opening_metrics(opening, versions[0] in {"2.6.0", "2.6.1", "2.7.0"}),
        "depth": {
            "rawAverage": _as_float(engaged["B8"].value, "raw average depth"),
            "postTemplateAverage": _as_float(engaged["C8"].value, "post-template average depth"),
            "netValidAverage": _as_float(net["C8"].value, "net-valid average depth"),
            "rawMedian": _as_float(engaged["B9"].value, "raw median depth"),
            "postTemplateMedian": _as_float(engaged["C9"].value, "post-template median depth"),
            "netValidMedian": _as_float(net["C9"].value, "net-valid median depth"),
            "rawFivePlusShare": _as_float(engaged["B10"].value, "raw five-plus share"),
            "postTemplateFivePlusShare": _as_float(engaged["C10"].value, "post-template five-plus share"),
            "netValidFivePlusShare": _as_float(net["C10"].value, "net-valid five-plus share"),
            "netValidSessions": _as_int(net["C6"].value, "net-valid sessions"),
        },
        "riskSummary": {
            "qualityCandidates": _as_int(overview["J5"].value, "quality candidates"),
            "qualityRate": _nullable_float(overview["K5"].value, "quality rate"),
            "scoreableAiReplies": _number_from_label(overview["L5"].value, "scoreable AI replies"),
            "aiSafetyCandidates": _as_int(overview["J6"].value, "AI safety candidates"),
            "aiSafetyRate": _nullable_float(overview["K6"].value, "AI safety rate"),
            "userRiskCandidates": _as_int(overview["J7"].value, "user risk candidates"),
            "userRiskRate": _nullable_float(overview["K7"].value, "user risk rate"),
            "reviewableUserInputs": _number_from_label(overview["L7"].value, "reviewable user inputs"),
            "candidateUsers": _as_int(overview["J8"].value, "candidate users"),
            "candidateUserRate": _nullable_float(overview["K8"].value, "candidate user rate"),
            "reviewableUsers": _number_from_label(overview["L8"].value, "reviewable users"),
            "priority": {"p0": priority_counts[0], "p1": priority_counts[1], "p2": priority_counts[2]},
            "aiParseFailures": _as_int(overview["J10"].value, "AI parse failures"),
        },
        "reopen": {
            "triggerSessions": _as_int(raw_ending["B18"].value, "reopen trigger sessions"),
            "triggerUsers": _as_int(raw_ending["D18"].value, "reopen trigger users"),
            "fiveMinuteSessions": _as_int(raw_ending["F18"].value, "five-minute sessions"),
            "fiveMinuteUsers": _as_int(raw_ending["H18"].value, "five-minute users"),
            "sessionRate": _as_float(raw_ending["B19"].value, "reopen session rate"),
            "userRate": _as_float(raw_ending["D19"].value, "reopen user rate"),
            "windowSeconds": _as_int(raw_ending["F19"].value, "reopen window"),
            "delayBuckets": delay_buckets,
        },
        "coverage": {
            "intentParsed": parsed_intents,
            "intentMissing": missing_intents,
            "intentRate": parsed_intents / (parsed_intents + missing_intents),
            "profileUsers": _as_int(profile["D4"].value, "normal profile users"),
            "profileEligibleUsers": _as_int(profile["B4"].value, "profile eligible users"),
            "profileRate": _as_float(profile["F4"].value, "profile coverage rate"),
            "locationUsers": locations["normalUsers"],
            "locationEligibleUsers": locations["eligibleUsers"],
            "locationRate": locations["coverageRate"],
            "constellationUsers": constellations["normalUsers"],
            "constellationEligibleUsers": constellations["eligibleUsers"],
            "constellationRate": constellations["coverageRate"],
        },
        "endings": {
            "validCohortSessions": valid_ending_sessions,
            "rawCohortSessions": raw_ending_sessions,
            "valid": _ending_rows(valid_ending),
            "raw": _ending_rows(raw_ending),
            "validIntents": _ending_intents(valid_ending, valid_ending_sessions, "valid"),
            "rawIntents": _ending_intents(raw_ending, raw_ending_sessions, "raw"),
        },
        "qualityCategories": _problem_rows(problems, 5, 13),
        "aiSafetyCategories": _problem_rows(problems, 14, 20),
        "userRiskCategories": _problem_rows(problems, 21, 27),
        "intents": {
            "primary": _intent_rows(intent, "A", "B", "C", intent_header + 1, 200),
            "secondary": _intent_rows(intent, "D", "E", "F", intent_header + 1, 200),
        },
        "demographics": demographic_rows,
        "locations": locations,
        "constellations": constellations,
    }
    if valid_intent_comparison is not None:
        result["validSecondaryIntentComparison"] = valid_intent_comparison
    return result


def build_snapshot(
    workbook_path: Path | str,
    manifest_path: Path | str,
    data_date: date,
    *,
    full_hash_check: bool = True,
    primary_source: Path | str | None = None,
) -> dict[str, Any]:
    workbook = Path(workbook_path).resolve()
    manifest_file = Path(manifest_path).resolve()
    _require(workbook.is_file(), "workbook does not exist")
    _require(manifest_file.is_file(), "manifest does not exist")
    _require(workbook.parent == manifest_file.parent, "workbook and manifest must be in the same package directory")

    manifest = json.loads(manifest_file.read_text(encoding="utf-8-sig"))
    source_hash, workbook_hash, scene_id, product_version = _validate_manifest(
        manifest,
        workbook,
        full_hash_check,
        Path(primary_source).resolve() if primary_source is not None else None,
    )
    if primary_source is not None:
        try:
            source_summary = validate_csv_data_date(primary_source, data_date, scene_id)
        except (CsvDateContractError, UnicodeError, csv.Error) as exc:
            raise SnapshotContractError(str(exc)) from exc
        _require(source_summary["rowCount"] == manifest["sources"][0].get("rows"), "primary source row count does not match manifest")
    metrics = _extract_workbook(workbook, (manifest["coreVersion"], manifest["rulesVersion"]))

    supplemental = None
    if manifest["coreVersion"] in {"2.6.0", "2.6.1", "2.7.0"}:
        entries = [item for item in manifest["outputs"] if item.get("role") == "supplemental_metrics"]
        _require(len(entries) == 1, "new metric contract requires supplemental metrics")
        item = entries[0]
        _require(Path(item["file"]).name == item["file"], "unsafe supplemental filename")
        companion = manifest_file.parent / item["file"]
        _require(sha256_file(companion) == item["sha256"], "supplemental hash mismatch")
        supplemental = json.loads(companion.read_text(encoding="utf-8-sig"))
        _require(supplemental.get("schemaVersion") == "teeni-base-supplemental-metrics/1.0.0", "unsupported supplemental schema")
        _require(supplemental.get("coreVersion") == manifest["coreVersion"] and supplemental.get("rulesVersion") == manifest["rulesVersion"], "supplemental version mismatch")
        _require(str(supplemental.get("sceneId")) == scene_id, "supplemental scene mismatch")
        metrics["supplemental"] = supplemental["primary"]

    payload = {
        "schemaVersion": SNAPSHOT_SCHEMA if manifest["coreVersion"] == "2.7.0" else PREVIOUS_PRODUCT_SNAPSHOT_SCHEMA,
        "dataDate": data_date.isoformat(),
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "productVersion": product_version,
        "sceneId": scene_id,
        "coreVersion": manifest["coreVersion"],
        "rulesVersion": manifest["rulesVersion"],
        "contractVersion": EXPECTED_CONTRACT,
        "sourceSha256": source_hash,
        "workbookSha256": workbook_hash,
        "metrics": metrics,
    }
    detail = next(item for item in manifest["outputs"] if item.get("role") == "primary_detail")
    payload["baseDetailSha256"] = detail["sha256"]
    payload["baseManifestSha256"] = sha256_file(manifest_file)
    if manifest["coreVersion"] == "2.7.0":
        entries = [item for item in manifest["outputs"] if item.get("role") == "session_structure"]
        _require(len(entries) == 1, "new contract requires session structure companion")
        item = entries[0]
        _require(Path(item["file"]).name == item["file"], "unsafe session structure filename")
        companion = manifest_file.parent / item["file"]
        _require(sha256_file(companion) == item["sha256"], "session structure hash mismatch")
        structure = json.loads(companion.read_text(encoding="utf-8-sig"))
        _require(structure.get("coreVersion") == manifest["coreVersion"] and structure.get("rulesVersion") == manifest["rulesVersion"] and str(structure.get("sceneId")) == scene_id, "session structure source mismatch")
        computed = from_detail(manifest_file.parent / detail["file"], scene_id, data_date)
        _require(structure.get("primary") == computed, "session structure independent reconciliation failed")
        metrics["sessionStructure"] = computed
        payload["sessionStructureSource"] = {"sourceSha256": source_hash, "detailSha256": detail["sha256"], "manifestSha256": payload["baseManifestSha256"]}
    validate_snapshot_payload(payload)
    assert_aggregate_only(payload)
    return payload


def _validate_profile_section(value: Any, label: str, expected_labels: list[str] | None = None) -> None:
    _require(isinstance(value, dict), f"snapshot {label} metrics are invalid")
    eligible = _as_int(value.get("eligibleUsers"), f"snapshot {label} eligible users")
    normal = _as_int(value.get("normalUsers"), f"snapshot {label} normal users")
    coverage = _as_float(value.get("coverageRate"), f"snapshot {label} coverage")
    _require(0 <= normal <= eligible, f"snapshot {label} user counts are invalid")
    _require(abs(coverage - normal / max(1, eligible)) < 1e-12, f"snapshot {label} coverage does not match counts")
    statuses = value.get("statuses")
    items = value.get("items")
    _require(isinstance(statuses, list) and isinstance(items, list), f"snapshot {label} rows are invalid")
    status_counts: dict[str, int] = {}
    for index, row in enumerate(statuses):
        _require(isinstance(row, dict), f"snapshot {label} status row {index} is invalid")
        status = _as_label(row.get("status"), f"snapshot {label} status row {index}")
        users = _as_int(row.get("users"), f"snapshot {label} status users row {index}")
        _require(users >= 0 and status not in status_counts, f"snapshot {label} status row {index} is invalid")
        status_counts[status] = users
    _require(sum(status_counts.values()) == eligible, f"snapshot {label} statuses do not cover users")
    _require(status_counts.get("正常", 0) == normal, f"snapshot {label} normal status mismatch")
    labels: list[str] = []
    item_users = 0
    for index, row in enumerate(items):
        _require(isinstance(row, dict), f"snapshot {label} item row {index} is invalid")
        item_label = _as_label(row.get("label"), f"snapshot {label} label row {index}")
        _require(item_label not in labels, f"snapshot {label} labels must be unique")
        labels.append(item_label)
        users = _as_int(row.get("users"), f"snapshot {label} users row {index}")
        sessions = _as_int(row.get("sessions"), f"snapshot {label} sessions row {index}")
        turns = _as_int(row.get("netValidTurns"), f"snapshot {label} turns row {index}")
        five_plus = _as_int(row.get("fivePlusSessions"), f"snapshot {label} five-plus row {index}")
        user_share = _as_float(row.get("userShare"), f"snapshot {label} user share row {index}")
        average = _as_float(row.get("averageDepth"), f"snapshot {label} depth row {index}")
        five_share = _as_float(row.get("fivePlusShare"), f"snapshot {label} five-plus share row {index}")
        sample_status = _as_label(row.get("sampleStatus"), f"snapshot {label} sample row {index}")
        _require(min(users, sessions, turns, five_plus) >= 0 and five_plus <= sessions, f"snapshot {label} counts row {index} are invalid")
        _require(abs(user_share - users / max(1, normal)) < 1e-12, f"snapshot {label} user share row {index} mismatch")
        _require(abs(average - turns / max(1, sessions)) < 1e-12, f"snapshot {label} average row {index} mismatch")
        _require(abs(five_share - five_plus / max(1, sessions)) < 1e-12, f"snapshot {label} five-plus row {index} mismatch")
        _require(sample_status == ("可描述" if users >= 30 else "样本不足"), f"snapshot {label} sample row {index} mismatch")
        item_users += users
    _require(item_users == normal, f"snapshot {label} items do not partition normal users")
    if expected_labels is not None:
        _require(labels == expected_labels, f"snapshot {label} categories are incomplete")


def _validate_ending_intents_payload(value: Any, cohort_sessions: int, label: str) -> None:
    _require(isinstance(value, dict), f"snapshot {label} ending intents are invalid")
    for kind in ("primary", "secondary"):
        rows = value.get(kind)
        _require(isinstance(rows, list) and bool(rows), f"snapshot {label} ending {kind} intents are missing")
        labels: set[str] = set()
        total = 0
        for index, row in enumerate(rows):
            _require(isinstance(row, dict), f"snapshot {label} ending {kind} row {index} is invalid")
            intent_label = _as_label(row.get("label"), f"snapshot {label} ending {kind} label {index}")
            _require(intent_label not in labels, f"snapshot {label} ending {kind} labels must be unique")
            labels.add(intent_label)
            count = _as_int(row.get("count"), f"snapshot {label} ending {kind} count {index}")
            share = _as_float(row.get("share"), f"snapshot {label} ending {kind} share {index}")
            _require(count >= 0, f"snapshot {label} ending {kind} count {index} must be non-negative")
            _require(abs(share - count / cohort_sessions) < 1e-12, f"snapshot {label} ending {kind} share {index} mismatch")
            total += count
        _require(total == cohort_sessions, f"snapshot {label} ending {kind} intent count mismatch")


def _validate_valid_secondary_intent_comparison(value: Any) -> None:
    _require(isinstance(value, dict), "snapshot valid secondary intent comparison is invalid")
    all_total = _as_int(value.get("allNetValidTurns"), "snapshot valid secondary intent all total")
    ending_total = _as_int(value.get("endingSessions"), "snapshot valid secondary intent ending total")
    rows = value.get("rows")
    _require(all_total > 0 and ending_total > 0, "snapshot valid secondary intent totals must be positive")
    _require(isinstance(rows, list) and bool(rows), "snapshot valid secondary intent rows are missing")
    labels: set[str] = set()
    summed_all = 0
    summed_ending = 0
    for index, row in enumerate(rows):
        _require(isinstance(row, dict), f"snapshot valid secondary intent row {index} is invalid")
        label = _as_label(row.get("label"), f"snapshot valid secondary intent label {index}")
        _require(label not in labels and label != "未返回", f"snapshot valid secondary intent label {index} is invalid")
        labels.add(label)
        all_count = _as_int(row.get("allCount"), f"snapshot valid secondary intent all count {index}")
        ending_count = _as_int(row.get("endingCount"), f"snapshot valid secondary intent ending count {index}")
        all_share = _as_float(row.get("allShare"), f"snapshot valid secondary intent all share {index}")
        ending_share = _as_float(row.get("endingShare"), f"snapshot valid secondary intent ending share {index}")
        share_delta = _as_float(row.get("shareDelta"), f"snapshot valid secondary intent delta {index}")
        _require(min(all_count, ending_count) >= 0, f"snapshot valid secondary intent counts {index} must be non-negative")
        _require(abs(all_share - all_count / all_total) < 1e-12, f"snapshot valid secondary intent all share {index} mismatch")
        _require(abs(ending_share - ending_count / ending_total) < 1e-12, f"snapshot valid secondary intent ending share {index} mismatch")
        _require(abs(share_delta - (ending_share - all_share)) < 1e-12, f"snapshot valid secondary intent delta {index} mismatch")
        summed_all += all_count
        summed_ending += ending_count
    _require(summed_all == all_total, "snapshot valid secondary intent all total mismatch")
    _require(summed_ending == ending_total, "snapshot valid secondary intent ending total mismatch")


def _validate_supplemental_values(value: Any) -> None:
    if isinstance(value, dict):
        if "numerator" in value or "denominator" in value:
            _require(set(value) == {"numerator", "denominator", "rate"}, "invalid ratio fields")
            n, d, rate = value["numerator"], value["denominator"], value["rate"]
            _require(type(n) is int and type(d) is int and 0 <= n <= d, "invalid ratio counts")
            _require(rate is None if d == 0 else isinstance(rate, (int, float)) and abs(rate - n / d) < 1e-12, "invalid ratio value")
        for child in value.values():
            _validate_supplemental_values(child)
    elif isinstance(value, list):
        for child in value:
            _validate_supplemental_values(child)
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        _require(math.isfinite(value) and value >= 0, "invalid supplemental number")


def _validate_risk_summary(risk: Any) -> None:
    _require(isinstance(risk, dict), "invalid risk summary")
    for rate_key, numerator_key, denominator_key in (
        ("qualityRate", "qualityCandidates", "scoreableAiReplies"),
        ("aiSafetyRate", "aiSafetyCandidates", "scoreableAiReplies"),
        ("userRiskRate", "userRiskCandidates", "reviewableUserInputs"),
        ("candidateUserRate", "candidateUsers", "reviewableUsers"),
    ):
        numerator, denominator = risk.get(numerator_key), risk.get(denominator_key)
        _require(type(numerator) is int and type(denominator) is int and numerator >= 0 and denominator >= 0, "invalid risk counts")
        if rate_key != "aiSafetyRate":
            _require(numerator <= denominator, "unique risk count exceeds population")
        actual = risk.get(rate_key)
        _require(actual is None and numerator == 0 if denominator == 0 else type(actual) in (int, float) and math.isfinite(actual) and abs(actual - numerator / denominator) < 1e-12, "risk rate does not match counts or no-sample state")


def validate_snapshot_payload(payload: Any) -> None:
    _require(isinstance(payload, dict), "snapshot must be an object")
    schema = payload.get("schemaVersion")
    _require(
        schema in (LEGACY_SNAPSHOT_SCHEMA, PREVIOUS_SNAPSHOT_SCHEMA, PREVIOUS_PRODUCT_SNAPSHOT_SCHEMA, SNAPSHOT_SCHEMA),
        "unsupported snapshot schema",
    )
    scene_id = str(payload.get("sceneId"))
    product_version = payload.get("productVersion")
    if schema in (PREVIOUS_PRODUCT_SNAPSHOT_SCHEMA, SNAPSHOT_SCHEMA):
        _require(product_version in PRODUCT_VERSION_SCENES, "snapshot product version is invalid")
        _require(
            PRODUCT_VERSION_SCENES[product_version] == scene_id,
            "snapshot product version does not match scene",
        )
        expected_versions = (EXPECTED_CORE, EXPECTED_RULES, EXPECTED_CONTRACT)
        if payload.get("rulesVersion") == "12.0.0":
            expected_versions = ("2.5.0", "12.0.0", EXPECTED_CONTRACT)
        elif payload.get("rulesVersion") == "14.0.0":
            expected_versions = ("2.6.0", "14.0.0", EXPECTED_CONTRACT)
        elif payload.get("rulesVersion") == "15.0.0":
            expected_versions = ("2.6.1", "15.0.0", EXPECTED_CONTRACT)
        elif payload.get("rulesVersion") == "16.0.0":
            expected_versions = ("2.7.0", "16.0.0", EXPECTED_CONTRACT)
    elif schema == PREVIOUS_SNAPSHOT_SCHEMA:
        _require(scene_id == "488", "previous snapshot scene must be 488")
        _require(product_version in (None, "M1"), "previous snapshot product version must be M1")
        expected_versions = (PREVIOUS_CORE, PREVIOUS_RULES, PREVIOUS_CONTRACT)
    else:
        _require(scene_id == "488", "legacy snapshot scene must be 488")
        _require(product_version in (None, "M1"), "legacy snapshot product version must be M1")
        expected_versions = (LEGACY_CORE, LEGACY_RULES, LEGACY_CONTRACT)
    _require(payload.get("coreVersion") == expected_versions[0], "snapshot core version mismatch")
    _require(payload.get("rulesVersion") == expected_versions[1], "snapshot rules version mismatch")
    _require(payload.get("contractVersion") == expected_versions[2], "snapshot contract mismatch")
    try:
        data_date = date.fromisoformat(payload.get("dataDate", ""))
    except (TypeError, ValueError) as exc:
        raise SnapshotContractError("snapshot data date is invalid") from exc
    if product_version == "M2":
        _require(data_date >= M2_FIRST_DATA_DATE, "M2 snapshots start on 2026-09-01")
    for key in ("sourceSha256", "workbookSha256"):
        value = payload.get(key)
        _require(isinstance(value, str) and re.fullmatch(r"[0-9a-f]{64}", value) is not None, f"{key} is invalid")
    _require(isinstance(payload.get("metrics"), dict), "snapshot metrics are missing")
    if payload.get("coreVersion") in {"2.6.0", "2.6.1", "2.7.0"}:
        extra = payload["metrics"].get("supplemental")
        _require(isinstance(extra, dict), "new contract requires supplemental metrics")
        _require(set(extra) == {"quality", "openingFunnel", "earlyExperience", "continuationProxies", "fallbackReopen", "observation", "inputQuality", "productModels"}, "unexpected supplemental sections")
        _validate_supplemental_values(extra)
    for section in (
        "volume",
        "participation",
        "opening",
        "depth",
        "riskSummary",
        "reopen",
        "coverage",
        "endings",
        "qualityCategories",
        "aiSafetyCategories",
        "userRiskCategories",
        "intents",
        "demographics",
        *(() if schema == LEGACY_SNAPSHOT_SCHEMA else ("locations", "constellations")),
    ):
        _require(section in payload["metrics"], f"snapshot section is missing: {section}")

    opening = payload["metrics"]["opening"]
    _require(isinstance(opening, dict), "snapshot opening metrics are invalid")
    exposures = _as_int(opening.get("exposures"), "snapshot opening exposures")
    opened = _as_int(opening.get("opened"), "snapshot opening opened")
    rate = _nullable_float(opening.get("rate"), "snapshot opening rate")
    _require(exposures >= 0, "snapshot opening exposures must be nonnegative")
    _require(0 <= opened <= exposures, "snapshot opening opened is invalid")
    _require(rate is None if not exposures else rate is not None and abs(rate - opened / exposures) < 1e-12, "snapshot opening rate does not match counts")
    structure = payload["metrics"].get("sessionStructure")
    _require(payload.get("coreVersion") != "2.7.0" or structure is not None, "new snapshot requires session structure")
    if structure is not None:
        try:
            validate_structure(structure)
        except ValueError as exc:
            raise SnapshotContractError(str(exc)) from exc
        volume = payload["metrics"]["volume"]
        _require((structure["totalTurns"], structure["totalSessions"]) == (volume["turns"], volume["sessions"]), "session structure volume mismatch")
        provenance = payload.get("sessionStructureSource", {})
        _require(set(provenance) == {"sourceSha256", "detailSha256", "manifestSha256"}, "session structure provenance missing")
        _require(all(isinstance(value, str) and re.fullmatch(r"[0-9a-f]{64}", value) for value in provenance.values()), "session structure provenance invalid")
        _require(provenance["sourceSha256"] == payload["sourceSha256"], "session structure raw source mismatch")
        if payload.get("baseDetailSha256"):
            _require(provenance["detailSha256"] == payload["baseDetailSha256"], "session structure detail source mismatch")
        if payload.get("baseManifestSha256"):
            _require(provenance["manifestSha256"] == payload["baseManifestSha256"], "session structure manifest source mismatch")
    if payload.get("coreVersion") == "2.7.0":
        _validate_risk_summary(payload["metrics"]["riskSummary"])
    if "openingCohorts" in payload["metrics"]:
        from .opening_cohorts import validate_opening_cohorts
        addon = payload["metrics"]["openingCohorts"]
        _require(structure is not None, "opening cohorts require session structure")
        try:
            validate_opening_cohorts(addon, structure)
        except (ValueError, TypeError, KeyError) as exc:
            raise SnapshotContractError(f"invalid opening cohorts: {exc}") from exc
        for key in ("productVersion", "dataDate", "sceneId"):
            _require(addon[key] == payload.get(key), f"opening cohort {key} mismatch")
        for source_key, payload_key in (("sourceSha256", "sourceSha256"), ("workbookSha256", "workbookSha256"),
                                        ("detailSha256", "baseDetailSha256"), ("manifestSha256", "baseManifestSha256")):
            _require(addon["source"][source_key] == payload.get(payload_key), f"opening cohort {source_key} mismatch")
    if schema != LEGACY_SNAPSHOT_SCHEMA:
        metrics = payload["metrics"]
        endings = metrics["endings"]
        _require(isinstance(endings, dict), "snapshot ending metrics are invalid")
        has_valid_intents = "validIntents" in endings
        has_raw_intents = "rawIntents" in endings
        _require(has_valid_intents == has_raw_intents, "snapshot ending intent cohorts must be published together")
        if has_valid_intents:
            valid_sessions = _as_int(endings.get("validCohortSessions"), "snapshot valid ending cohort")
            raw_sessions = _as_int(endings.get("rawCohortSessions"), "snapshot raw ending cohort")
            _require(valid_sessions > 0 and raw_sessions > 0, "snapshot ending cohort counts must be positive")
            _validate_ending_intents_payload(endings["validIntents"], valid_sessions, "valid")
            _validate_ending_intents_payload(endings["rawIntents"], raw_sessions, "raw")
        if "validSecondaryIntentComparison" in metrics:
            comparison = metrics["validSecondaryIntentComparison"]
            _validate_valid_secondary_intent_comparison(comparison)
            if has_valid_intents:
                ending_rows = {
                    row["label"]: row["count"]
                    for row in endings["validIntents"]["secondary"]
                    if row["label"] != "未返回" and row["count"] != 0
                }
                comparison_rows = {
                    row["label"]: row["endingCount"]
                    for row in comparison["rows"]
                    if row["endingCount"] != 0
                }
                _require(comparison_rows == ending_rows, "snapshot valid secondary intent ending rows mismatch")
        _validate_profile_section(metrics["locations"], "locations")
        _validate_profile_section(metrics["constellations"], "constellations", [
            "摩羯座", "水瓶座", "双鱼座", "白羊座", "金牛座", "双子座",
            "巨蟹座", "狮子座", "处女座", "天秤座", "天蝎座", "射手座",
        ])
        coverage = metrics["coverage"]
        for prefix, section in (("location", metrics["locations"]), ("constellation", metrics["constellations"])):
            _require(_as_int(coverage.get(f"{prefix}Users"), f"snapshot {prefix} users") == section["normalUsers"], f"snapshot {prefix} users mismatch")
            _require(_as_int(coverage.get(f"{prefix}EligibleUsers"), f"snapshot {prefix} eligible users") == section["eligibleUsers"], f"snapshot {prefix} eligible users mismatch")
            _require(abs(_as_float(coverage.get(f"{prefix}Rate"), f"snapshot {prefix} rate") - section["coverageRate"]) < 1e-12, f"snapshot {prefix} coverage mismatch")
